"""Build the final report (.docx).

Every number in the text and tables is read from project artefacts at build time:
backend/data/benchmark_results.json (copied from the executed notebook), the notebook's own
EDA output, and docs/report/figures/*.json written by make_figures.py. Run from the repo root:

    backend/.venv/Scripts/python.exe docs/report/make_figures.py      # charts, diagrams, equations
    node docs/report/capture_screens.mjs docs/report/screens <profile> # demo screenshots
    backend/.venv/Scripts/python.exe docs/report/build_report.py
"""
import json
import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from PIL import Image

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))
from docx_kit import GRAY, NAVY, SLATE, Report, _set_font  # noqa: E402

FIG, EQ, SCREENS = HERE / "figures", HERE / "eq", HERE / "screens"
OUT = ROOT / "docs" / "Bao_cao_KG4RS_RippleNet_CKAN.docx"

BENCH = json.loads((ROOT / "backend" / "data" / "benchmark_results.json").read_text(encoding="utf-8"))
DS = BENCH["datasets"]
KG_STATS = json.loads((FIG / "kg_structure.json").read_text(encoding="utf-8"))
NOTEBOOK = BENCH["source"]
NAMES = {"movie": "MovieLens-20M", "book": "Book-Crossing", "music": "Last.FM"}
MODELS = ["MostPopular", "MF", "RippleNet", "CKAN"]

# Notebook cell 12 output (EDA summary table), copied verbatim. Per-user figures count positives only.
EDA = {
    "movie": dict(users=2500, items=16954, ratings=247606, positives=123803, sparsity=99.71, median=25.0, avg=49.5, cold=5.4, top20=83.4, entities=102569, relations=32, triples=499474, branching=4.87),
    "book": dict(users=17860, items=14967, ratings=139746, positives=69873, sparsity=99.97, median=1.0, avg=3.9, cold=88.3, top20=68.8, entities=77903, relations=25, triples=151500, branching=1.94),
    "music": dict(users=1872, items=3846, ratings=42346, positives=21173, sparsity=99.71, median=11.0, avg=11.3, cold=5.6, top20=76.0, entities=9366, relations=60, triples=15518, branching=1.66),
}
# Notebook CONFIG cell (CKAN and MF) and RIPPLE_CONFIG (RippleNet).
CONFIG = {
    "movie": dict(n_layer=1, utss=32, itss=64, batch=2048, epochs=10),
    "book": dict(n_layer=2, utss=16, itss=64, batch=1024, epochs=8),
    "music": dict(n_layer=2, utss=8, itss=64, batch=1024, epochs=10),
}


def vn(number, digits=0):
    """1234567.8 -> '1.234.567,8' (Vietnamese separators)."""
    text = f"{number:,.{digits}f}"
    return text.replace(",", "_").replace(".", ",").replace("_", ".")


def f4(value):
    return f"{value:.4f}".replace(".", ",")


RIPPLE_CONFIG = {
    "movie": dict(dim=16, n_hop=2, n_memory=32, lr="0,02", kge="0,01", l2="10⁻⁷", batch=1024, epochs=10),
    "book": dict(dim=4, n_hop=2, n_memory=32, lr="0,001", kge="0,01", l2="10⁻⁵", batch=1024, epochs=10),
    "music": dict(dim=16, n_hop=2, n_memory=32, lr="0,02", kge="0,01", l2="10⁻⁵", batch=1024, epochs=10),
}


def gain(d, i):
    sp = DS[d]["sparsity"]["auc"]
    return (sp["CKAN"][i] - sp["MF"][i]) / sp["MF"][i] * 100


def signed(value):
    return f"{'+' if value >= 0 else '−'}{abs(value):.1f}".replace(".", ",") + "%"


def screen(name, max_width=1900, crop=None):
    """JPEG copy of a screenshot, small enough to keep the document light.
    `crop` keeps only the top part: height = crop x width."""
    out = SCREENS / "jpg"
    out.mkdir(exist_ok=True)
    target = out / f"{name}.jpg"
    with Image.open(SCREENS / f"{name}.png") as im:
        im = im.convert("RGB")
        if crop and im.height > im.width * crop:
            im = im.crop((0, 0, im.width, int(im.width * crop)))
        if im.width > max_width:
            im = im.resize((max_width, int(im.height * max_width / im.width)), Image.LANCZOS)
        im.save(target, quality=90, optimize=True)
    return target


R = Report()
doc = R.doc
H1, H2, H3 = R.h1, R.h2, R.h3
CENTER = WD_ALIGN_PARAGRAPH.CENTER
_pending_equations = []


def eq(key, scale=1.0):
    """Reserve an equation number now; the equation itself is placed after the next paragraph
    or list, so the sentence that introduces it always comes first."""
    label = R.equation_label()
    _pending_equations.append((key, label, scale))
    return label


def _flush_equations():
    while _pending_equations:
        key, label, scale = _pending_equations.pop(0)
        R.equation(EQ / f"{key}.png", label, scale)


def P(text, **kwargs):
    if _pending_equations:
        kwargs["keep"] = True
    paragraph = R.p(text, **kwargs)
    _flush_equations()
    return paragraph


def B(items, **kwargs):
    R.bullets(items, **kwargs)
    _flush_equations()


def fig(name, label, caption, width=15.5):
    R.figure(FIG / f"{name}.png", label, caption, width)


# Screenshot widths (cm), chosen so a screenshot and a few lines of text share a page.
WIDE = 15.5


def shot(name, label, caption, width=16.0, crop=None):
    R.figure(screen(name, crop=crop), label, caption, width)


# =============================================================================== cover
def cover():
    """Cover laid out like the faculty's reference report: double-line page frame, school and
    faculty, logo, report and course name, topic, member table, class, supervisor, place and date."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def line(text, size, after=6, bold=True):
        return P(text, indent=False, align=CENTER, size=size, bold=bold, color="000000", after=after)

    def blank(n=1):
        for _ in range(n):
            P("", indent=False, align=CENTER, after=0)

    line("TRƯỜNG ĐẠI HỌC BÁCH KHOA", 16, after=2)
    line("KHOA CÔNG NGHỆ THÔNG TIN", 16, after=10)
    logo = doc.add_paragraph()
    logo.alignment = CENTER
    logo.add_run().add_picture(str(FIG / "cover_logo.jpeg"), width=Cm(2.6))
    blank(2)
    line("BÁO CÁO", 20, after=4)
    line("[TÊN HỌC PHẦN]", 20, after=14)
    line("Nghiên cứu đồ thị tri thức cho hệ gợi ý với RippleNet và CKAN", 20, after=6)
    blank(2)

    members = [("[Họ và tên sinh viên 1]", "[MSSV]"), ("[Họ và tên sinh viên 2]", "[MSSV]"), ("[Họ và tên sinh viên 3]", "[MSSV]")]
    table = doc.add_table(rows=1 + len(members), cols=2)
    table.alignment = 1
    header = table.rows[0].cells[0].merge(table.rows[0].cells[1])
    header.paragraphs[0].alignment = CENTER
    _set_font(header.paragraphs[0].add_run("SINH VIÊN THỰC HIỆN"), size=13, bold=True)
    for r, (name, student_id) in enumerate(members, 1):
        for c, (text, width) in enumerate(((name, 5.6), (student_id, 3.8))):
            cell = table.rows[r].cells[c]
            cell.width = Cm(width)
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT if c == 0 else CENTER
            para.paragraph_format.space_after = Pt(3)
            _set_font(para.add_run(text), size=13)
    for row in table.rows:
        for cell in row.cells:
            cell.paragraphs[0].paragraph_format.space_after = Pt(3)

    blank(1)
    line("Lớp học phần: [Mã lớp học phần]", 13, after=4)
    line("GIẢNG VIÊN HƯỚNG DẪN: [Học hàm, học vị. Họ và tên]", 13, after=4)
    blank(3)
    line("Đà Nẵng, 10/2026", 13)

    # Frame on the first page only, same line style and colour as the reference report.
    sect = doc.sections[0]._sectPr
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:display"), "firstPage")
    for side, style, space in (("top", "thinThickMediumGap", "2"), ("left", "thinThickMediumGap", "4"),
                               ("bottom", "thickThinMediumGap", "1"), ("right", "thickThinMediumGap", "4")):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), style)
        el.set(qn("w:sz"), "36")
        el.set(qn("w:space"), space)
        el.set(qn("w:color"), "156082")
        borders.append(el)
    sect.find(qn("w:pgMar")).addnext(borders)


cover()
R.doc.sections[0].different_first_page_header_footer = True
R.page_numbers()

# =============================================================================== front matter
H1("MỞ ĐẦU", numbered=False)
P("Hệ gợi ý đã trở thành thành phần cốt lõi của các nền tảng trực tuyến: từ xem phim, nghe nhạc đến mua sách, người dùng ngày càng dựa vào danh sách được cá nhân hoá để tìm ra thứ mình cần giữa hàng chục nghìn lựa chọn. Phương pháp kinh điển cho bài toán này là lọc cộng tác, tức dự đoán sở thích của một người từ hành vi của những người giống họ. Lọc cộng tác hoạt động tốt khi dữ liệu tương tác dày, nhưng suy giảm rõ rệt khi ma trận tương tác thưa và gần như bất lực trước người dùng hay sản phẩm mới, hiện tượng thường gọi là khởi động lạnh.")
P("Một hướng khắc phục được nghiên cứu mạnh trong những năm gần đây là đưa đồ thị tri thức (Knowledge Graph, KG) vào hệ gợi ý. Đồ thị tri thức mô tả sản phẩm bằng các bộ ba có cấu trúc, ví dụ một bộ phim nối với đạo diễn, diễn viên, thể loại của nó. Nhờ đó hai sản phẩm chưa từng được cùng một người dùng tương tác vẫn có thể liên hệ với nhau qua một thực thể chung, và mỗi gợi ý có thể được lý giải bằng một đường đi cụ thể trên đồ thị.")
P("Đề tài này nghiên cứu hai mô hình tiêu biểu của hướng lan truyền trên đồ thị tri thức: **RippleNet** (CIKM 2018), mô hình đầu tiên lan truyền sở thích người dùng theo từng lớp như gợn sóng, và **CKAN** (SIGIR 2020), mô hình kế thừa ý tưởng đó nhưng lan truyền ở cả hai phía người dùng và sản phẩm, đồng thời mã hoá tường minh tín hiệu cộng tác. Hai mô hình được cài đặt lại, huấn luyện và so sánh với hai phương pháp cơ sở trên ba tập dữ liệu chuẩn thuộc ba miền khác nhau là MovieLens-20M, Book-Crossing và Last.FM. Từ mô hình CKAN đã huấn luyện, nhóm xây dựng một hệ thống web minh hoạ cho phép xem gợi ý, xem đường dẫn tri thức lý giải từng gợi ý và đối chiếu trực tiếp các kết quả thực nghiệm.")
P("Báo cáo gồm tám chương. Chương 1 giới thiệu đề tài. Chương 2 trình bày cơ sở lý thuyết. Chương 3 và Chương 4 phân tích chi tiết RippleNet và CKAN. Chương 5 mô tả dữ liệu. Chương 6 trình bày thực nghiệm và kết quả. Chương 7 mô tả hệ thống minh hoạ. Chương 8 kết luận và nêu hướng phát triển.")

H1("MỤC LỤC", numbered=False)
R.toc('TOC \\o "1-3" \\h \\z \\u')
H1("DANH MỤC HÌNH ẢNH", numbered=False)
R.toc('TOC \\h \\z \\t "Chú thích hình,1"')
H1("DANH MỤC BẢNG", numbered=False)
R.toc('TOC \\h \\z \\t "Chú thích bảng,1"')

# =============================================================================== chapter 1
H1("GIỚI THIỆU ĐỀ TÀI")
H2("Bối cảnh và lý do chọn đề tài")
P("Khi mở một ứng dụng xem phim hay nghe nhạc, thứ người dùng thấy đầu tiên thường là một danh sách gợi ý. Đằng sau danh sách đó phần lớn là lọc cộng tác: hệ thống tìm những người có lịch sử giống bạn rồi gợi ý những thứ họ đã thích. Cách làm này đơn giản và hiệu quả, nhưng có một điểm yếu rõ ràng là cần rất nhiều dữ liệu tương tác.")
P(f"Trên thực tế dữ liệu hiếm khi đủ. Ở cả ba tập dữ liệu của đề tài, hơn 99% số ô trong bảng người dùng và sản phẩm là ô trống ({vn(EDA['movie']['sparsity'], 2)}% với phim, {vn(EDA['music']['sparsity'], 2)}% với nhạc, {vn(EDA['book']['sparsity'], 2)}% với sách). Riêng tập sách, {vn(EDA['book']['cold'], 1)}% người dùng có không quá 5 lượt tương tác dương. Với những người dùng như vậy, hệ thống gần như không có gì để so sánh.")
P("Đồ thị tri thức là một cách bù vào chỗ thiếu đó. Thay vì chỉ biết “người này đã xem phim kia”, hệ thống biết thêm bộ phim do ai đạo diễn, ai đóng, thuộc thể loại gì. Hai bộ phim chưa từng có chung một khán giả vẫn có thể được nối với nhau qua một đạo diễn chung. Thông tin này được kỳ vọng giúp gợi ý tốt hơn khi dữ liệu ít, và còn trả lời được câu hỏi người dùng hay đặt ra: vì sao tôi được gợi ý thứ này? Đề tài kiểm tra kỳ vọng đó bằng thực nghiệm.")
P("Nhóm chọn RippleNet và CKAN vì đây là hai mô hình nối tiếp nhau trong cùng một hướng. RippleNet là mô hình đầu tiên cho sở thích của người dùng lan dần trên đồ thị tri thức. CKAN ra đời hai năm sau và bổ sung những gì RippleNet còn thiếu. Đặt hai mô hình cạnh nhau cho thấy rõ mỗi cải tiến đem lại điều gì, và đó cũng là lý do nhóm cài đặt lại cả hai thay vì chỉ dẫn lại kết quả có sẵn.")

H2("Mục tiêu")
B([
    "Nắm vững cơ sở lý thuyết của hệ gợi ý dùng đồ thị tri thức, đặc biệt là nhóm phương pháp lan truyền.",
    "Phân tích chi tiết kiến trúc, công thức và hàm mất mát của RippleNet và CKAN; chỉ ra điểm giống và khác nhau.",
    "Cài đặt lại hai mô hình bằng PyTorch cùng hai phương pháp cơ sở (MostPopular và Matrix Factorization), huấn luyện và đánh giá trên ba tập dữ liệu thuộc ba miền.",
    "Đánh giá trên ba khía cạnh: dự đoán nhấp chuột (CTR), xếp hạng Top-K và khả năng chịu dữ liệu thưa.",
    "Xây dựng hệ thống web minh hoạ dùng mô hình CKAN đã huấn luyện, có lý giải từng gợi ý bằng đường dẫn tri thức.",
])
H2("Phạm vi")
B([
    "**Dữ liệu:** ba tập chuẩn đã được các nghiên cứu trước ghép sẵn với đồ thị tri thức: MovieLens-20M, Book-Crossing và Last.FM. Riêng tập phim được lấy mẫu 2.500 người dùng để phù hợp tài nguyên huấn luyện.",
    "**Mô hình:** MostPopular, Matrix Factorization, RippleNet và CKAN. Đề tài không huấn luyện lại embedding đồ thị tri thức riêng (TransE, TransR) và không xét các mô hình mạng nơ-ron đồ thị khác như KGCN hay KGAT ngoài phần tổng quan.",
    "**Phản hồi:** phản hồi ngầm dạng nhị phân (có hoặc không tương tác).",
    "**Hệ thống:** ứng dụng minh hoạ chạy cục bộ, phục vụ trình diễn và kiểm chứng kết quả, chưa hướng tới triển khai sản phẩm.",
])
H2("Phương pháp nghiên cứu")
B([
    "**Nghiên cứu tài liệu:** đọc hai bài báo gốc và mã nguồn công bố kèm theo của các tác giả để hiểu đúng mô hình và quy trình tiền xử lý.",
    "**Cài đặt và thực nghiệm:** viết lại các mô hình trong một notebook duy nhất theo mã nguồn của tác giả, dùng chung hàm dựng tập bộ ba, chung cách chia dữ liệu 6:2:2 và chung quy trình chọn mô hình theo AUC trên tập kiểm định.",
    "**Đánh giá định lượng:** ROC-AUC, F1, Accuracy cho bài toán CTR; Recall@K cho bài toán xếp hạng; ROC-AUC theo tỷ lệ dữ liệu huấn luyện cho bài toán độ thưa.",
    "**Kiểm chứng bằng hệ thống:** triển khai mô hình đã huấn luyện vào một ứng dụng web và đối chiếu rằng hệ thống tái tạo đúng embedding do quá trình thực nghiệm xuất ra.",
], numbered=False)

# =============================================================================== chapter 2
H1("CƠ SỞ LÝ THUYẾT")
H2("Bài toán gợi ý với phản hồi ngầm")
P("Xét tập người dùng $U$ và tập sản phẩm (item) $V$. Ma trận tương tác $Y$ có phần tử $y_{uv}=1$ nếu người dùng $u$ đã tương tác với item $v$ (xem, nghe, đánh giá cao) và $y_{uv}=0$ trong trường hợp ngược lại. Giá trị 0 không có nghĩa là người dùng không thích mà chỉ là chưa có tương tác được ghi nhận. Bài toán là học một hàm dự đoán xác suất người dùng $u$ sẽ tương tác với item $v$ mà họ chưa từng gặp.")
e_mf = eq("mf")
P(f"Phương pháp phân rã ma trận (Matrix Factorization, MF) gán cho mỗi người dùng và mỗi item một vector ẩn $d$ chiều và dự đoán bằng tích vô hướng của chúng cộng với các hệ số chệch, như công thức {e_mf}. Mỗi vector chỉ được cập nhật khi người dùng hoặc item tương ứng xuất hiện trong dữ liệu huấn luyện, vì vậy MF phụ thuộc hoàn toàn vào mật độ của $Y$.", keep=True)

H2("Dữ liệu thưa và khởi động lạnh")
P("Hai vấn đề gắn liền với lọc cộng tác là dữ liệu thưa và khởi động lạnh. Dữ liệu thưa nghĩa là mỗi người dùng chỉ tương tác với một phần rất nhỏ danh mục, khiến hai người dùng bất kỳ hiếm khi có item chung để so sánh. Khởi động lạnh là trường hợp cực đoan: người dùng mới hoặc item mới chưa có tương tác nào. Cách xử lý phổ biến là bổ sung thông tin phụ (side information) như thuộc tính sản phẩm, mạng xã hội hoặc, như trong đề tài này, đồ thị tri thức.")

H2("Đồ thị tri thức")
e_kg = eq("kg")
P(f"Đồ thị tri thức là một đồ thị có hướng, đa quan hệ, biểu diễn tri thức dưới dạng các bộ ba (đầu, quan hệ, đuôi) như công thức {e_kg}, trong đó $E$ là tập thực thể và $R$ là tập quan hệ. Ví dụ bộ ba (Forrest Gump, đạo diễn, Robert Zemeckis) cho biết thực thể đuôi là đạo diễn của thực thể đầu. Trong hệ gợi ý, mỗi item được ánh xạ tới một thực thể của đồ thị; tập các cặp ánh xạ này ký hiệu là $A$.")
e_task = eq("task")
P(f"Bài toán gợi ý có đồ thị tri thức được phát biểu như công thức {e_task}: hàm dự đoán $F$ với tham số Θ nhận thêm đồ thị $G$ làm đầu vào bên cạnh ma trận tương tác $Y$.")
f = R.fig()
P(f"{f} minh hoạ vai trò của đồ thị tri thức bằng một đồ thị con thật lấy từ hệ thống của đề tài. Người dùng #1 đã thích hai bộ phim *Field of Dreams* và *The Truman Show*. Cả hai có chung nhà thiết kế bối cảnh Nancy Haigh với *Forrest Gump*, nên bộ phim này được nối với sở thích của người dùng qua một thực thể trung gian, dù chưa hề có người dùng nào được dùng để so sánh.")
fig("diagram_kg_example", f, "Đồ thị con thật nối các phim người dùng đã thích với một phim được gợi ý qua thực thể chung", 15.5)

H2("Các hướng tiếp cận hệ gợi ý dùng đồ thị tri thức")
f = R.fig()
P(f"Các phương pháp KG4RS thường được chia thành ba nhóm như {f}.")
fig("diagram_taxonomy", f, "Ba nhóm phương pháp đưa đồ thị tri thức vào hệ gợi ý", 15.5)
B([
    "**Dựa trên embedding.** Học trước vector cho thực thể và quan hệ bằng các mô hình như TransE, TransR rồi dùng chúng làm đặc trưng cho mô hình gợi ý (CKE, DKN). Nhóm này linh hoạt nhưng mục tiêu học embedding phục vụ bài toán hoàn thiện đồ thị nhiều hơn là gợi ý.",
    "**Dựa trên đường dẫn.** Khai thác các mẫu kết nối giữa người dùng và item (meta-path, meta-graph) như PER. Nhóm này trực quan nhưng phụ thuộc vào việc thiết kế đường dẫn thủ công, khó tối ưu và khó áp dụng sang miền dữ liệu khác.",
    "**Dựa trên lan truyền.** Kết hợp ưu điểm của hai nhóm trên: tự động khám phá các đường dẫn bằng cách lan truyền thông tin qua nhiều bước láng giềng và học đầu-cuối cùng mục tiêu gợi ý. RippleNet là mô hình mở đầu nhóm này; KGCN, KGAT và CKAN là các bước phát triển tiếp theo.",
])

H2("Cơ chế attention")
P("Khi một thực thể có nhiều láng giềng, không phải láng giềng nào cũng quan trọng như nhau đối với một dự đoán cụ thể. Cơ chế attention gán cho mỗi láng giềng một trọng số không âm, chuẩn hoá bằng hàm softmax để tổng bằng 1, rồi lấy tổng có trọng số của các vector láng giềng. Cả RippleNet và CKAN đều dùng attention để tổng hợp các bộ ba trong một lớp lan truyền; hai mô hình khác nhau ở cách tính trọng số, sẽ được phân tích ở Chương 3 và Chương 4.")

H2("Các độ đo đánh giá")
e_auc = eq("auc")
P(f"**ROC-AUC** đo xác suất mô hình chấm một mẫu dương cao hơn một mẫu âm chọn ngẫu nhiên, theo công thức {e_auc} với $P$ và $N$ là tập mẫu dương và mẫu âm trong tập kiểm thử. AUC không phụ thuộc ngưỡng quyết định nên phù hợp để so sánh chất lượng xếp thứ tự của các mô hình.", keep=True)
e_f1 = eq("f1")
P(f"**F1 và Accuracy** (công thức {e_f1}) được tính sau khi nhị phân hoá điểm dự đoán tại ngưỡng 0,5. Hai độ đo này phản ánh thêm việc điểm số có được hiệu chỉnh quanh ngưỡng hay không.", keep=True)
e_rec = eq("recall")
P(f"**Recall@K** (công thức {e_rec}) đo tỷ lệ item thật sự liên quan trong tập kiểm thử của một người dùng xuất hiện trong K gợi ý đầu, sau khi đã loại các item người dùng tương tác trong tập huấn luyện. Đây là độ đo sát với kịch bản sử dụng thực tế nhất.", keep=True)

# =============================================================================== chapter 3
H1("MÔ HÌNH RIPPLENET")
P("RippleNet được Hongwei Wang và các cộng sự công bố tại hội nghị CIKM 2018 [1]. Đây là mô hình đầu tiên thuộc nhóm lan truyền: nó đưa đồ thị tri thức vào hệ gợi ý theo cách học đầu-cuối, không cần embedding học trước và không cần thiết kế meta-path.")
H2("Ý tưởng")
f = R.fig()
P(f"Tên gọi của mô hình xuất phát từ hình ảnh gợn sóng trên mặt nước. Mỗi item người dùng từng tương tác giống một viên sỏi rơi xuống mặt hồ: nó kích hoạt các thực thể kề nó trên đồ thị, các thực thể đó lại kích hoạt các thực thể xa hơn. Sở thích của người dùng lan ra thành từng lớp, càng xa điểm xuất phát thì tín hiệu càng yếu ({f}). Nhiều gợn sóng xuất phát từ nhiều item trong lịch sử chồng lên nhau tạo thành phân bố sở thích của người dùng đối với một item ứng viên.")
fig("diagram_ripple_sets", f, "Sở thích người dùng lan truyền theo từng bước trên đồ thị tri thức", 15.5)

H2("Tập thực thể liên quan và ripple set")
e1 = eq("rn_entity", 0.86)
P(f"Với người dùng $u$, gọi $V_{{u}}$ là tập item họ đã tương tác. Tập thực thể liên quan bậc $k$ được định nghĩa đệ quy theo công thức {e1}: thực thể bậc $k$ là các thực thể đuôi của những bộ ba có thực thể đầu thuộc bậc $k-1$. Các item trong lịch sử đóng vai trò hạt giống (bậc 0).")
e2 = eq("rn_set")
P(f"Ripple set bậc $k$ của người dùng là tập các bộ ba xuất phát từ tập thực thể bậc $k-1$, theo công thức {e2}. Trong thực tế kích thước ripple set tăng rất nhanh theo $k$, nên mô hình lấy mẫu một số lượng cố định bộ ba ở mỗi bước, và số bước $H$ thường chỉ là 1 đến 3: đi quá xa thì nhiễu nhiều hơn tín hiệu.")

H2("Lan truyền sở thích")
f = R.fig()
P(f"Kiến trúc của RippleNet được trình bày trong {f}. Đầu vào là một cặp người dùng và item ứng viên; đầu ra là xác suất người dùng nhấp vào item đó.")
fig("diagram_ripplenet", f, "Kiến trúc RippleNet: mỗi bước lan truyền sinh một vector phản hồi, các phản hồi cộng lại thành biểu diễn người dùng", 15.5)
e3 = eq("rn_p", 0.9)
P(f"Mỗi item $v$ có một embedding $d$ chiều. Ở bước lan truyền đầu tiên, mỗi bộ ba trong ripple set bậc 1 được gán một xác suất liên quan $p_{{i}}$ bằng cách so sánh item ứng viên với thực thể đầu trong không gian của quan hệ, theo công thức {e3}. Ở đây quan hệ được biểu diễn bằng một ma trận $d×d$ và thực thể đầu bằng một vector. Ý nghĩa của $p_{{i}}$ là: xét theo quan hệ này, item ứng viên giống thực thể đầu đến mức nào. Chẳng hạn hai bộ phim có thể rất giống nhau khi xét theo đạo diễn nhưng ít giống khi xét theo thể loại, và việc đưa quan hệ vào phép so sánh cho phép mô hình phân biệt hai trường hợp đó.")
e4 = eq("rn_o")
P(f"Sau đó các thực thể đuôi được cộng lại với trọng số là xác suất liên quan, cho ra vector phản hồi bậc 1 của người dùng đối với item $v$ (công thức {e4}). Vector này có thể hiểu là: trong những thứ kề với lịch sử của người dùng, phần nào liên quan đến item đang xét.")
P("Lan truyền tiếp tục theo cùng cách ở các bước sau: ở bước 2, vector phản hồi bậc 1 thay cho embedding của item để tính xác suất liên quan trên ripple set bậc 2, và cứ thế đến bước $H$. Nhờ vậy sở thích của người dùng được truyền dần ra các thực thể xa hơn nhưng vẫn luôn được điều hướng bởi item ứng viên.")
e5 = eq("rn_u")
e6 = eq("rn_y")
P(f"Biểu diễn cuối cùng của người dùng là tổng các vector phản hồi của mọi bước (công thức {e5}). Xác suất nhấp chuột được tính bằng hàm sigmoid của tích vô hướng giữa biểu diễn người dùng và embedding item (công thức {e6}). Cần lưu ý rằng trong RippleNet, biểu diễn người dùng phụ thuộc vào item ứng viên: cùng một người dùng sẽ có vector khác nhau khi được xét với các item khác nhau.")

H2("Hàm mất mát")
e7 = eq("rn_loss", 0.78)
P(f"RippleNet được huấn luyện bằng cách cực đại hoá xác suất hậu nghiệm của tham số mô hình khi biết ma trận tương tác và đồ thị tri thức. Hàm mất mát thu được có ba thành phần, như công thức {e7}:")
B([
    "**Mất mát cross-entropy** giữa nhãn tương tác thật và xác suất dự đoán, trên các mẫu dương và các mẫu âm được lấy mẫu.",
    "**Mất mát tái tạo đồ thị tri thức**, đo sai lệch giữa lát cắt $I_{r}$ của tensor chỉ thị quan hệ và ma trận tái tạo từ embedding thực thể và ma trận quan hệ. Thành phần này buộc embedding phải giữ được cấu trúc của đồ thị.",
    "**Chuẩn hoá L2** trên embedding item, embedding thực thể và ma trận quan hệ để hạn chế quá khớp.",
])

H2("Khả năng lý giải")
P("Vì xác suất liên quan $p_{i}$ được tính cho từng bộ ba, có thể truy lại những đường đi có trọng số lớn nhất từ lịch sử của người dùng đến item được gợi ý và dùng chúng làm lời giải thích. Bài báo gốc minh hoạ điều này bằng các đường đi kiểu “người dùng đã xem phim A, phim A có diễn viên X, diễn viên X đóng trong phim B”. Hệ thống minh hoạ của đề tài áp dụng cùng nguyên lý, trình bày ở Chương 7.")

H2("Ưu điểm và hạn chế")
B([
    "**Ưu điểm.** Tự động khám phá đường dẫn, không cần meta-path thủ công; học đầu-cuối cùng mục tiêu gợi ý; có khả năng lý giải; giảm ảnh hưởng của dữ liệu thưa vì người dùng được biểu diễn qua các thực thể trên đồ thị.",
    "**Hạn chế thứ nhất: item là vector tĩnh.** Chỉ phía người dùng được lan truyền; item ứng viên chỉ là một embedding tra bảng, không tận dụng láng giềng của chính nó trên đồ thị.",
    "**Hạn chế thứ hai: tín hiệu cộng tác không được mã hoá tường minh.** Mô hình chỉ đi theo các cạnh của đồ thị tri thức, không dùng thông tin “những người cùng tương tác với item này còn tương tác với gì”.",
    "**Hạn chế thứ ba: quan hệ tốn tham số.** Mỗi quan hệ là một ma trận $d×d$, và trọng số attention là một dạng song tuyến tính đơn giản, khó nắm bắt các tương tác phi tuyến giữa thực thể đầu và quan hệ.",
    "**Hạn chế thứ tư: kích thước ripple set.** Số bộ ba tăng nhanh theo số bước, buộc phải lấy mẫu; chất lượng phụ thuộc vào kích thước mẫu và số bước được chọn.",
])

H2("Cài đặt RippleNet trong đề tài")
t = R.tab()
P(f"Bản cài đặt trong notebook là bản chuyển sang PyTorch của mã nguồn TensorFlow do tác giả công bố [15], giữ nguyên cách biểu diễn, hàm mất mát và cấu hình theo từng tập dữ liệu ({t}). Mã nguồn gốc không có cấu hình cho Last.FM; tập nhạc dùng cấu hình của tập phim với hệ số L2 của tập sách.", keep=True)
R.table(t, "Cấu hình RippleNet theo từng tập dữ liệu (lấy từ notebook)", ["Tập dữ liệu", "Số chiều d", "Số bước H", "Bộ ba mỗi bước", "Tốc độ học", "λ KGE", "λ L2", "Batch", "Epoch"],
        [[NAMES[d], c["dim"], c["n_hop"], c["n_memory"], c["lr"], c["kge"], c["l2"], vn(c["batch"]), c["epochs"]] for d, c in RIPPLE_CONFIG.items()],
        [3.0, 1.6, 1.5, 1.9, 1.7, 1.4, 1.4, 1.5, 1.4], size=10)
e8 = eq("rn_impl_p", 0.86)
e9 = eq("rn_impl_loss", 0.8)
P("Mã nguồn của tác giả có vài chi tiết không nêu trong bài báo, và bản cài đặt làm theo mã nguồn:")
B([
    "**Quan hệ là ma trận $d×d$**, đúng như bài báo. Vì số tham số của quan hệ tăng theo bình phương $d$, tác giả dùng số chiều nhỏ: 16 cho tập phim và 4 cho tập sách.",
    f"**Embedding item được cập nhật sau mỗi bước.** Thay vì dùng vector phản hồi bậc trước để tính attention ở bước sau, mã nguồn cộng vector phản hồi vào embedding item rồi nhân với một ma trận biến đổi học được (chế độ plus_transform). Điểm dự đoán là tích vô hướng của embedding item sau bước cuối với tổng các vector phản hồi (công thức {e8}).",
    f"**Hàm mất mát có đủ ba thành phần** (công thức {e9}): cross-entropy, thành phần đồ thị tri thức thưởng cho các bộ ba trong ripple set có điểm song tuyến tính giữa thực thể đầu, quan hệ và thực thể đuôi cao, và chuẩn hoá L2 trên embedding của các bộ ba đó.",
    "**Ripple set mỗi bước có 32 bộ ba**, lấy mẫu ngẫu nhiên đều từ các bộ ba xuất phát từ tập thực thể của bước trước, có lặp lại nếu không đủ.",
])
P("Sau mỗi epoch mô hình được đánh giá trên tập kiểm định, và trạng thái có AUC kiểm định cao nhất được giữ lại.")

# =============================================================================== chapter 4
H1("MÔ HÌNH CKAN")
P("CKAN (Collaborative Knowledge-aware Attentive Network) được Ze Wang, Guangyan Lin, Huobin Tan, Qinghong Chen và Xiyang Liu công bố tại hội nghị SIGIR 2020 [2]. Đây là mô hình trọng tâm của đề tài: nó được phân tích kỹ nhất, đạt AUC cao nhất trên hai trong ba tập dữ liệu và là mô hình chạy trong hệ thống minh hoạ.")

H2("Động cơ")
P("Các tác giả CKAN xuất phát từ một nhận xét về những mô hình lan truyền có trước như RippleNet và KGCN: chúng tập trung vào việc mã hoá các liên kết tri thức trong đồ thị mà không làm nổi bật tín hiệu cộng tác ẩn trong tương tác giữa người dùng và item. Tín hiệu cộng tác ở đây là thông tin kiểu “những người thích item này cũng thích item kia”, vốn là nền tảng của lọc cộng tác.")
P("Một cách khác để đưa tín hiệu cộng tác vào là gộp tương tác và đồ thị tri thức thành một đồ thị thống nhất rồi lan truyền trên đó, như KGAT. Tuy nhiên cách này coi quan hệ “tương tác” ngang hàng với các quan hệ tri thức, và việc lan truyền nhiều bước qua các cạnh tương tác dễ dẫn tới những item không liên quan. Bài báo nêu ví dụ bộ phim *Forrest Gump* (hài, chính kịch) lan tới *12 Monkeys* (bí ẩn, khoa học viễn tưởng) chỉ sau ba bước tương tác.")
P("CKAN đề xuất một cách kết hợp tự nhiên hơn dựa trên hai thiết kế:")
B([
    "**Lan truyền dị thể** (heterogeneous propagation), gồm lan truyền cộng tác để mã hoá tường minh tín hiệu cộng tác vào tập thực thể khởi đầu, và lan truyền tri thức để mở rộng tập đó trên đồ thị.",
    "**Embedding có chú ý nhận biết tri thức** (knowledge-aware attentive embedding), gán trọng số khác nhau cho các thực thể đuôi tuỳ theo thực thể đầu và quan hệ dẫn tới chúng.",
])

H2("Tổng quan kiến trúc")
f = R.fig()
P(f"{f} trình bày kiến trúc CKAN. Khác với RippleNet, mô hình có hai nhánh đối xứng: nhánh người dùng và nhánh item đều đi qua bốn khối giống nhau là lan truyền cộng tác, lan truyền tri thức, lớp attention và bộ tổng hợp. Hai nhánh dùng chung bảng embedding thực thể, embedding quan hệ và trọng số của mạng attention. Điểm dự đoán là sigmoid của tích vô hướng giữa hai biểu diễn cuối.")
fig("diagram_ckan", f, "Kiến trúc hai nhánh của CKAN: người dùng và item đều được lan truyền trên đồ thị tri thức", 15.5)

H2("Lan truyền dị thể")
H3("Lan truyền cộng tác")
e1 = eq("ck_eu0")
P(f"CKAN không dùng một vector ẩn độc lập cho mỗi người dùng. Thay vào đó, người dùng được biểu diễn bằng chính các item họ đã tương tác: tập thực thể khởi đầu của người dùng gồm các thực thể ứng với những item trong lịch sử của họ (công thức {e1}, với φ là phép ánh xạ item sang thực thể).")
e2 = eq("ck_ev0", 0.86)
P(f"Tương tự, một item được đặc trưng bởi những item khác được cùng người dùng tương tác. Gọi $V_{{v}}$ là tập các item mà ít nhất một người dùng đã tương tác đồng thời với $v$. Tập thực thể khởi đầu của item gồm các thực thể ứng với $V_{{v}}$, hợp với thực thể của chính $v$ (công thức {e2}). Việc giữ lại thực thể gốc của $v$ giúp thông tin của item ban đầu luôn được nhấn mạnh và giảm sai lệch do lan truyền nhiều tầng.")
f = R.fig()
P(f"{f} minh hoạ hai tập khởi đầu này. Đây chính là chỗ tín hiệu cộng tác đi vào mô hình: hai item chưa có cạnh nào nối nhau trên đồ thị tri thức vẫn chia sẻ thông tin nếu chúng có chung người dùng.")
fig("diagram_collab", f, "Lan truyền cộng tác: tập thực thể khởi đầu của người dùng và của item", 15.0)

H3("Lan truyền tri thức")
e3 = eq("ck_prop", 0.8)
P(f"Từ tập khởi đầu, CKAN mở rộng dần trên đồ thị tri thức theo đúng tinh thần ripple set của RippleNet. Với $o$ là người dùng hoặc item, tập thực thể tầng $l$ và tập bộ ba tầng $l$ được định nghĩa đệ quy như công thức {e3}. Các thực thể kề nhau trên đồ thị thường liên hệ chặt chẽ, nên các tập mở rộng này làm giàu biểu diễn của cả người dùng lẫn item. Số tầng $L$ là siêu tham số; mỗi tầng lấy mẫu một số bộ ba cố định.")
P("Điểm khác biệt căn bản so với RippleNet là lan truyền tri thức được thực hiện cho cả hai phía. Item ứng viên không còn là một vector tra bảng mà cũng được mô tả bằng các thực thể quanh nó.")

H2("Embedding có chú ý nhận biết tri thức")
P("Một thực thể đuôi mang ý nghĩa khác nhau tuỳ theo nó được nối tới từ thực thể đầu nào và qua quan hệ gì. Bài báo lấy ví dụ: *Forrest Gump* và *Cast Away* giống nhau nhiều khi xét theo đạo diễn hoặc diễn viên, nhưng ít giống khi xét theo thể loại hoặc biên kịch. Lấy trung bình đều các thực thể đuôi sẽ làm mất sự phân biệt đó.")
e4 = eq("ck_a")
P(f"CKAN vì vậy gán cho mỗi bộ ba thứ $i$ trong tập bộ ba tầng $l$ một trọng số π phụ thuộc vào embedding thực thể đầu và embedding quan hệ, rồi nhân với embedding thực thể đuôi (công thức {e4}).")
e5 = eq("ck_mlp", 0.82)
P(f"Hàm π được hiện thực bằng một mạng nơ-ron truyền thẳng ba tầng nhận đầu vào là phép nối embedding thực thể đầu và embedding quan hệ, với hàm kích hoạt ReLU ở hai tầng đầu và sigmoid ở tầng cuối (công thức {e5}).")
e6 = eq("ck_soft", 0.9)
e7 = eq("ck_layer")
P(f"Các trọng số được chuẩn hoá bằng softmax trên toàn bộ tập bộ ba của tầng (công thức {e6}). Biểu diễn tầng $l$ của người dùng hoặc item là tổng có trọng số của các embedding thực thể đuôi (công thức {e7}).")
f = R.fig()
P(f"{f} tóm tắt đường đi của dữ liệu qua lớp attention.")
fig("diagram_attention", f, "Lớp attention nhận biết tri thức của CKAN", 15.5)
P("So với attention của RippleNet, có hai khác biệt đáng chú ý. Thứ nhất, trọng số của CKAN chỉ phụ thuộc vào bộ ba, không phụ thuộc vào đối tượng ở phía bên kia; nhờ đó biểu diễn item không phụ thuộc người dùng và có thể tính trước, một tính chất được hệ thống minh hoạ tận dụng (Chương 7). Thứ hai, trọng số được tính bằng một mạng phi tuyến thay cho dạng song tuyến tính, và quan hệ chỉ cần một vector thay cho một ma trận.")
e8 = eq("ck_zero")
P(f"Ngoài các tầng lan truyền, tập thực thể khởi đầu cũng được đưa vào biểu diễn vì nó gần đối tượng gốc nhất và chứa tín hiệu cộng tác. Biểu diễn tầng 0 là trung bình embedding của các thực thể trong tập khởi đầu (công thức {e8}). Với item, embedding của thực thể gốc ứng với chính item đó cũng được giữ lại thành một thành phần riêng.")

H2("Bộ tổng hợp")
e9 = eq("ck_sets", 0.9)
P(f"Sau lan truyền, mỗi phía có một tập các vector biểu diễn theo từng tầng (công thức {e9}); tập của item có thêm thành phần gốc. Bộ tổng hợp gộp các vector này thành một biểu diễn duy nhất. Bài báo khảo sát ba lựa chọn:")
e10 = eq("ck_agg_sum", 0.82)
e11 = eq("ck_agg_cat", 0.9)
B([
    f"**Tổng (sum)**: cộng các vector sau một phép biến đổi tuyến tính rồi qua hàm kích hoạt phi tuyến; và **pooling**: lấy giá trị lớn nhất theo từng chiều (công thức {e10}).",
    f"**Nối (concat)**: nối các vector của mọi tầng rồi qua phép biến đổi tuyến tính (công thức {e11}). Cách này giữ riêng thông tin của từng tầng thay vì trộn lẫn.",
])
P("Thực nghiệm trong bài báo cho thấy bộ tổng hợp nối cho kết quả tốt nhất, và đây cũng là lựa chọn của đề tài.")

H2("Dự đoán và huấn luyện")
e12 = eq("ck_y")
e13 = eq("ck_loss", 0.84)
P(f"Điểm dự đoán là sigmoid của tích vô hướng giữa biểu diễn người dùng và biểu diễn item (công thức {e12}). Mô hình được huấn luyện bằng hàm mất mát cross-entropy trên các mẫu dương và các mẫu âm được lấy mẫu với số lượng bằng nhau cho mỗi người dùng, cộng với chuẩn hoá L2 trên toàn bộ tham số (công thức {e13}, trong đó $J$ là mất mát cross-entropy và $P$ là phân bố lấy mẫu âm). Tối ưu bằng Adam.")

H2("Các siêu tham số chính")
B([
    "**Số tầng lan truyền $L$.** Tăng $L$ cho phép nhìn xa hơn trên đồ thị nhưng cũng đưa thêm nhiễu. Bài báo dùng 1 tầng cho MovieLens-20M và 2 tầng cho Book-Crossing và Last.FM.",
    "**Kích thước tập bộ ba** của người dùng và của item ở mỗi tầng. Tập lớn chứa nhiều thông tin hơn nhưng tốn bộ nhớ và có thể loãng tín hiệu.",
    "**Số chiều embedding $d$**, dùng chung cho thực thể và quan hệ.",
    "**Bộ tổng hợp** (sum, pooling hoặc concat).",
])

H2("So sánh RippleNet và CKAN")
t = R.tab()
P(f"{t} tổng hợp các khác biệt chính giữa hai mô hình.", keep=True)
R.table(t, "So sánh RippleNet và CKAN", ["Tiêu chí", "RippleNet (CIKM 2018)", "CKAN (SIGIR 2020)"], [
    ["Hướng lan truyền", "Một phía: chỉ người dùng", "Hai phía: người dùng và item"],
    ["Biểu diễn item", "Vector tra bảng, tĩnh", "Tổng hợp từ thực thể gốc và các tầng láng giềng"],
    ["Tín hiệu cộng tác", "Không mã hoá tường minh", "Mã hoá qua tập thực thể khởi đầu (lan truyền cộng tác)"],
    ["Trọng số attention", "Song tuyến tính, phụ thuộc item ứng viên", "Mạng ba tầng trên (thực thể đầu, quan hệ), không phụ thuộc phía bên kia"],
    ["Biểu diễn quan hệ", "Ma trận d × d", "Vector d chiều"],
    ["Gộp các tầng", "Cộng các vector phản hồi", "Sum, pooling hoặc concat (concat tốt nhất)"],
    ["Hàm mất mát", "Cross-entropy + tái tạo KG + L2", "Cross-entropy + L2"],
    ["Tính trước biểu diễn item", "Không cần (item là vector tra bảng)", "Được, vì không phụ thuộc người dùng"],
], [3.6, 5.6, 6.8], left_cols=(0, 1, 2))

H2("Cài đặt CKAN trong đề tài")
t = R.tab()
P(f"Lớp mô hình trong notebook giữ nguyên mã nguồn PyTorch do nhóm tác giả công bố [16]; khi nạp cùng một bộ trọng số, hai bản cho ra điểm dự đoán trùng nhau. Cấu hình theo từng tập dữ liệu được nêu ở {t}. Số chiều embedding là 64, tốc độ học 0,002, hệ số suy giảm trọng số 10⁻⁵ và bộ tổng hợp nối được dùng chung cho cả ba tập.", keep=True)
R.table(t, "Cấu hình CKAN theo từng tập dữ liệu (lấy từ notebook)", ["Tập dữ liệu", "Số tầng L", "Bộ ba người dùng mỗi tầng", "Bộ ba item mỗi tầng", "Batch size", "Số epoch"],
        [[NAMES[d], c["n_layer"], c["utss"], c["itss"], vn(c["batch"]), c["epochs"]] for d, c in CONFIG.items()],
        [3.4, 2.0, 3.2, 2.9, 2.2, 2.3])
e14 = eq("ck_impl", 0.84)
P(f"Biểu diễn cuối và điểm dự đoán của bản cài đặt được cho bởi công thức {e14}. Mã nguồn của tác giả có vài chi tiết khác với mô tả trong bài báo, và bản cài đặt làm theo mã nguồn:")
B([
    "**Tập khởi đầu của item** gồm mọi item được cùng người dùng tương tác trong tập huấn luyện. Item gốc tự có mặt trong tập này vì nó nằm trong lịch sử của chính những người dùng đó. Item chưa có ai tương tác trong tập huấn luyện thì lấy chính nó làm tập khởi đầu.",
    "**Bộ tổng hợp nối không có phép biến đổi tuyến tính theo sau**; các vector tầng được nối trực tiếp. Ở phía item, thành phần đầu tiên là embedding của thực thể gốc; trung bình của tập khởi đầu chỉ được thêm vào khi dùng bộ tổng hợp tổng hoặc pooling.",
    "**Mạng attention không có hệ số chệch** ở cả ba tầng.",
    "**Mỗi tầng lấy mẫu một số bộ ba cố định**, ngẫu nhiên đều trên các bộ ba xuất phát từ tập thực thể của tầng trước, có lặp lại nếu không đủ. Tầng không có bộ ba nào thì lặp lại tầng trước.",
    "**Chọn mô hình theo tập kiểm định.** Sau mỗi epoch, mô hình được đánh giá trên tập kiểm định và trạng thái có AUC cao nhất được giữ lại để đánh giá trên tập kiểm thử.",
])

# =============================================================================== chapter 5
H1("DỮ LIỆU VÀ TIỀN XỬ LÝ")
H2("Ba tập dữ liệu")
t = R.tab()
P(f"Đề tài dùng ba tập dữ liệu công khai thuộc ba miền, cùng các đồ thị tri thức đã được các tác giả RippleNet và CKAN ghép sẵn. {t} tổng hợp các thống kê do notebook thực nghiệm tính ra.", keep=True)
R.table(t, "Thống kê ba tập dữ liệu sau tiền xử lý", ["Chỉ số", NAMES["movie"], NAMES["book"], NAMES["music"]], [
    ["Miền dữ liệu", "Phim", "Sách", "Âm nhạc (nghệ sĩ)"],
    ["Số người dùng"] + [vn(EDA[d]["users"]) for d in EDA],
    ["Số item"] + [vn(EDA[d]["items"]) for d in EDA],
    ["Số tương tác (gồm mẫu âm)"] + [vn(EDA[d]["ratings"]) for d in EDA],
    ["Số tương tác dương"] + [vn(EDA[d]["positives"]) for d in EDA],
    ["Độ thưa của ma trận (%)"] + [vn(EDA[d]["sparsity"], 2) for d in EDA],
    ["Số tương tác dương trung vị mỗi người dùng"] + [vn(EDA[d]["median"], 0) for d in EDA],
    ["Người dùng có ≤ 5 tương tác dương (%)"] + [vn(EDA[d]["cold"], 1) for d in EDA],
    ["Tỷ trọng tương tác dương của 20% item phổ biến nhất (%)"] + [vn(EDA[d]["top20"], 1) for d in EDA],
    ["Số thực thể trong KG"] + [vn(EDA[d]["entities"]) for d in EDA],
    ["Số loại quan hệ"] + [vn(EDA[d]["relations"]) for d in EDA],
    ["Số bộ ba"] + [vn(EDA[d]["triples"]) for d in EDA],
], [6.6, 3.1, 3.1, 3.2])
P(f"Các chỉ số theo người dùng và độ thưa chỉ tính trên tương tác dương, vì mẫu âm là do bước tiền xử lý sinh ra. Ba tập có đặc điểm rất khác nhau. Tập phim có tương tác dày nhất theo từng người dùng (trung vị {vn(EDA['movie']['median'], 0)} tương tác dương) và đồ thị tri thức lớn nhất. Tập sách có nhiều người dùng nhất nhưng mỗi người chỉ có trung vị {vn(EDA['book']['median'], 0)} tương tác dương, là tập khó nhất về độ thưa. Tập nhạc nhỏ nhất cả về tương tác lẫn đồ thị tri thức, nhưng có nhiều loại quan hệ nhất. Ở cả ba tập, tương tác tập trung mạnh vào nhóm item phổ biến: 20% item phổ biến nhất chiếm từ {vn(min(EDA[d]['top20'] for d in EDA), 1)}% đến {vn(max(EDA[d]['top20'] for d in EDA), 1)}% số tương tác dương.")

H2("Quy trình tiền xử lý")
f = R.fig()
P(f"Quy trình tiền xử lý ({f}) tuân theo mã nguồn gốc của hai bài báo để kết quả có thể đối chiếu.")
fig("diagram_data_pipeline", f, "Quy trình tiền xử lý dữ liệu tương tác", 15.5)
B([
    "**Ánh xạ item sang thực thể.** Chỉ giữ các item có thực thể tương ứng trong đồ thị tri thức; mỗi item được đánh số lại từ 0, các thực thể còn lại của đồ thị được đánh số tiếp sau.",
    "**Nhị phân hoá nhãn.** Với phim, đánh giá từ 4 sao trở lên được coi là tương tác dương. Với sách và nhạc, mọi tương tác được ghi nhận đều là dương vì dữ liệu vốn rất thưa.",
    "**Lấy mẫu âm.** Với mỗi người dùng, lấy ngẫu nhiên số item chưa tương tác bằng đúng số tương tác dương của họ làm mẫu âm (tỷ lệ 1:1).",
    "**Lấy mẫu người dùng tập phim.** Chọn ngẫu nhiên 2.500 người dùng (hạt giống ngẫu nhiên 42) để vừa bộ nhớ huấn luyện.",
])
P("Kết quả là hai mảng số nguyên cho mỗi tập: mảng tương tác gồm các bộ (người dùng, item, nhãn) và mảng đồ thị tri thức gồm các bộ (đầu, quan hệ, đuôi).")

H2("Khám phá dữ liệu tương tác")
f1, f2 = R.fig(), R.fig()
P(f"{f1} và {f2} là hai hình do notebook thực nghiệm sinh ra ở bước khám phá dữ liệu. Hình thứ nhất so sánh quy mô, độ thưa, tỷ lệ người dùng ít tương tác và phân bố mức độ hoạt động của người dùng. Hình thứ hai mô tả quy mô và cấu trúc của đồ thị tri thức.")
fig("nb_cell13_0", f1, f"Khám phá phân phối dữ liệu tương tác của ba tập (hình xuất từ notebook {NOTEBOOK})", 15.5)
fig("nb_cell13_1", f2, "Quy mô và cấu trúc đồ thị tri thức của ba tập (hình xuất từ notebook)", 15.5)

H2("Cấu trúc đồ thị tri thức")
f = R.fig()
P(f"{f} cho thấy các loại quan hệ chiếm nhiều bộ ba nhất ở từng tập. Ở tập phim, quan hệ diễn viên áp đảo. Ở tập sách, gần như toàn bộ đồ thị xoay quanh quan hệ tác giả và tác phẩm. Ở tập nhạc, nhiều bộ ba nói về các bộ phim mà nghệ sĩ góp mặt, loại thông tin ít liên quan trực tiếp đến gu âm nhạc.")
fig("chart_relations", f, "Tám loại quan hệ có nhiều bộ ba nhất ở mỗi đồ thị tri thức", 13.5)
f = R.fig()
m, tails, cover_ = KG_STATS["median_triples_per_item"], KG_STATS["tails_with_next_hop_pct"], KG_STATS["top10_with_kg_path_pct"]
P(f"Ba đồ thị còn khác nhau về hình dạng, điều ảnh hưởng trực tiếp đến lợi ích của việc lan truyền nhiều tầng ({f}). Mỗi phim có trung vị {vn(m[0])} bộ ba, trong khi mỗi cuốn sách và mỗi nghệ sĩ chỉ có trung vị {vn(m[1])} bộ ba. Ở tập sách, {vn(tails[1], 1)}% thực thể đuôi còn có cạnh đi tiếp nên tầng lan truyền thứ hai có nội dung; ở tập phim tỷ lệ này là {vn(tails[0], 1)}% và ở tập nhạc chỉ {vn(tails[2], 1)}%. Hệ quả là khả năng lý giải bằng đường dẫn cũng khác nhau: đo trên {KG_STATS['users_sampled']} người dùng ngẫu nhiên mỗi tập, {vn(cover_[0], 0)}% gợi ý top-10 của CKAN ở tập phim có ít nhất một đường dẫn qua thực thể chung với lịch sử người dùng, so với {vn(cover_[1], 0)}% ở tập sách và {vn(cover_[2], 0)}% ở tập nhạc.")
fig("chart_kg_structure", f, "Độ dày, độ sâu của đồ thị tri thức và tỷ lệ gợi ý có đường dẫn tri thức (đo trên hệ thống của đề tài)", 16.0)

# =============================================================================== chapter 6
H1("THỰC NGHIỆM VÀ KẾT QUẢ")
H2("Thiết lập thực nghiệm")
f = R.fig()
P(f"Toàn bộ thực nghiệm được thực hiện trong notebook *{NOTEBOOK}*, chạy trên GPU Tesla T4 với PyTorch 2.11. Quy trình chung được mô tả ở {f}.")
fig("diagram_experiment", f, "Quy trình thực nghiệm: chia dữ liệu, huấn luyện bốn mô hình và ba nhóm đánh giá", 15.5)
t = R.tab()
B([
    "**Chia dữ liệu.** Xáo trộn ngẫu nhiên (hạt giống 42) rồi chia 60% huấn luyện, 20% kiểm định, 20% kiểm thử. Người dùng không có tương tác dương nào trong tập huấn luyện bị loại khỏi cả ba tập, theo cách làm của mã nguồn CKAN gốc.",
    "**Huấn luyện.** MF và CKAN dùng số chiều embedding 64, bộ tối ưu Adam với tốc độ học 0,002 và suy giảm trọng số 10⁻⁵. RippleNet dùng cấu hình riêng của mã nguồn gốc (Bảng 3.1). Cấu hình CKAN theo từng tập đã nêu ở Bảng 4.2.",
    "**Chọn mô hình.** MF, RippleNet và CKAN được đánh giá trên tập kiểm định sau mỗi epoch; trạng thái có AUC kiểm định cao nhất được dùng để đo trên tập kiểm thử.",
    "**Xếp hạng Top-K.** Recall@K được tính theo cách của mã nguồn CKAN: chọn ngẫu nhiên 100 người dùng có mặt ở cả tập huấn luyện và tập kiểm thử; ứng viên là mọi item xuất hiện trong hai tập này, trừ các item người dùng đã gặp trong tập huấn luyện. Hạt giống chọn người dùng được cố định để bốn mô hình được đo trên cùng một nhóm.",
    "**Độ thưa.** Giữ lại 10%, 20%, 40%, 60%, 80% và 100% tập huấn luyện. Ở mỗi mức, lịch sử người dùng, tập khởi đầu và các tập bộ ba được dựng lại chỉ từ phần dữ liệu được giữ, rồi MF, RippleNet và CKAN được huấn luyện lại từ đầu. Để các mức so sánh được với nhau, AUC được đo trên cùng một nhóm người dùng: những người đã có ít nhất một tương tác dương ngay ở mức 10%.",
])
P(f"{t} cho biết kích thước các tập sau khi chia.", keep=True)
R.table(t, "Kích thước tập huấn luyện, kiểm định và kiểm thử", ["Tập dữ liệu", "Tổng tương tác", "Huấn luyện", "Kiểm định", "Kiểm thử"],
        [[NAMES[d], vn(DS[d]["split"]["ratings"]), vn(DS[d]["split"]["train"]), vn(DS[d]["split"]["eval"]), vn(DS[d]["split"]["test"])] for d in DS],
        [3.8, 3.2, 3.0, 3.0, 3.0])

H2("Các mô hình so sánh")
e_pop = eq("pop")
B([
    f"**MostPopular.** Chấm điểm mỗi item bằng số lượt tương tác dương của nó trong tập huấn luyện, chuẩn hoá theo item phổ biến nhất (công thức {e_pop}). Mọi người dùng nhận cùng một danh sách; đây là mức sàn để đo giá trị của cá nhân hoá.",
    "**Matrix Factorization (MF).** Đại diện cho lọc cộng tác thuần, không dùng đồ thị tri thức.",
    "**RippleNet.** Lan truyền một phía trên đồ thị tri thức (Chương 3).",
    "**CKAN.** Lan truyền hai phía với attention nhận biết tri thức (Chương 4).",
])

H2("Kết quả dự đoán CTR")
t, f = R.tab(), R.fig()
P(f"{t} và {f} trình bày kết quả trên tập kiểm thử. Giá trị tốt nhất của mỗi cột trong từng tập dữ liệu được in đậm.", keep=True)
rows, bold = [], set()
for d in DS:
    for metric_i, key in enumerate(("auc", "f1", "acc")):
        best = max(DS[d]["models"][m_][key] for m_ in MODELS)
        for m_i, m_ in enumerate(MODELS):
            if DS[d]["models"][m_][key] == best:
                bold.add((len(rows) + m_i, 2 + metric_i))
    for m_ in MODELS:
        r_ = DS[d]["models"][m_]
        rows.append([NAMES[d], m_, f4(r_["auc"]), f4(r_["f1"]), f4(r_["acc"])])
R.table(t, "Kết quả dự đoán CTR trên tập kiểm thử", ["Tập dữ liệu", "Mô hình", "ROC-AUC", "F1-Score", "Accuracy"], rows, [3.6, 3.4, 2.8, 2.8, 2.8], left_cols=(0, 1), bold_cells=bold)
fig("chart_ctr", f, "ROC-AUC, F1 và Accuracy của bốn mô hình trên ba tập dữ liệu", 16.0)
mv, bk, ms = DS["movie"]["models"], DS["book"]["models"], DS["music"]["models"]
P(f"**Tập sách và tập nhạc: CKAN dẫn đầu cả ba độ đo.** Ở tập sách, CKAN đạt AUC {f4(bk['CKAN']['auc'])}, cao hơn MF ({f4(bk['MF']['auc'])}) và RippleNet ({f4(bk['RippleNet']['auc'])}). Ở tập nhạc khoảng cách rõ hơn: {f4(ms['CKAN']['auc'])} so với {f4(ms['MF']['auc'])} của MF và {f4(ms['RippleNet']['auc'])} của RippleNet. CKAN cũng dẫn đầu về F1 và Accuracy trên hai tập này.")
P(f"**RippleNet không vượt được MF về AUC trên tập sách và tập nhạc.** Lan truyền một phía trên đồ thị tri thức, với item chỉ là một vector tra bảng, chưa đủ để thắng lọc cộng tác thuần ở hai tập này; phần hơn của CKAN vì thế đến từ những gì CKAN bổ sung so với RippleNet, tức lan truyền phía item và tập khởi đầu mang tín hiệu cộng tác. Cũng cần lưu ý RippleNet dùng số chiều nhỏ hơn nhiều (16 và 4 so với 64) theo cấu hình của mã nguồn gốc.")
P(f"**Tập phim: bốn mô hình gần như ngang nhau về AUC.** AUC dao động trong khoảng {f4(min(mv[m_]['auc'] for m_ in MODELS))} đến {f4(max(mv[m_]['auc'] for m_ in MODELS))}; MF nhỉnh hơn hai mô hình đồ thị về F1 và Accuracy. Khi mỗi người dùng đã có trung vị {vn(EDA['movie']['median'], 0)} tương tác dương, ma trận tương tác tự nó đủ thông tin và đồ thị tri thức không đem lại thêm lợi ích đo được.")
P(f"**MostPopular có AUC cao nhưng F1 rất thấp.** Ở tập phim, MostPopular đạt AUC {f4(mv['MostPopular']['auc'])}, cao nhất bảng, nhưng F1 chỉ {f4(mv['MostPopular']['f1'])}. Có hai nguyên nhân. Mẫu âm được lấy ngẫu nhiên đều trên danh mục nên phần lớn là item ít phổ biến, trong khi mẫu dương tập trung ở item phổ biến; chỉ riêng độ phổ biến đã tách khá tốt hai nhóm. Ngược lại, điểm của MostPopular là tần suất chuẩn hoá, hầu hết nhỏ hơn ngưỡng 0,5, nên sau khi nhị phân hoá mô hình gần như không dự đoán mẫu nào là dương. Vì vậy AUC của MostPopular cần được đọc cùng với F1 và Recall@K.")

H2("Kết quả xếp hạng Top-K")
t, f = R.tab(), R.fig()
P(f"{t} và {f} trình bày Recall@K.", keep=True)
rows, bold = [], set()
ks = ["5", "10", "20", "50", "100"]
for d in DS:
    for k_i, k in enumerate(ks):
        best = max(DS[d]["models"][m_]["recall"][k] for m_ in MODELS)
        for m_i, m_ in enumerate(MODELS):
            if DS[d]["models"][m_]["recall"][k] == best:
                bold.add((len(rows) + m_i, 2 + k_i))
    for m_ in MODELS:
        rows.append([NAMES[d], m_] + [f4(DS[d]["models"][m_]["recall"][k]) for k in ks])
R.table(t, "Recall@K trên 100 người dùng mẫu", ["Tập dữ liệu", "Mô hình"] + [f"R@{k}" for k in ks], rows, [3.2, 2.9, 1.9, 1.9, 1.9, 1.9, 2.0], left_cols=(0, 1), bold_cells=bold, size=10)
fig("chart_recall", f, "Recall@K của bốn mô hình trên ba tập dữ liệu", 16.0)
P(f"**Tập nhạc: CKAN tốt nhất từ K = 10 trở lên.** Recall@10 của CKAN là {f4(ms['CKAN']['recall']['10'])}, so với {f4(ms['MF']['recall']['10'])} của MF; ở K = 100 khoảng cách là {f4(ms['CKAN']['recall']['100'])} so với {f4(ms['MF']['recall']['100'])}. Riêng ở K = 5, MF nhỉnh hơn ({f4(ms['MF']['recall']['5'])} so với {f4(ms['CKAN']['recall']['5'])}). Đây là tập mà ưu thế của CKAN thể hiện ở cả CTR lẫn xếp hạng.")
P(f"**Tập phim và tập sách: các phương pháp không dùng đồ thị xếp hạng tốt hơn.** Ở tập phim, MF dẫn đầu ở mọi K (Recall@10 {f4(mv['MF']['recall']['10'])}), tiếp theo là MostPopular ({f4(mv['MostPopular']['recall']['10'])}) rồi CKAN ({f4(mv['CKAN']['recall']['10'])}). Ở tập sách, MostPopular dẫn đầu ở mọi K (Recall@10 {f4(bk['MostPopular']['recall']['10'])}), dù CKAN có AUC cao nhất. AUC cao trên các cặp đã lấy mẫu không tự động chuyển thành xếp hạng tốt trên toàn bộ danh mục. Một lý do là tương tác tập trung mạnh vào item phổ biến (Bảng 5.1), nên chỉ cần xếp đúng các item phổ biến lên đầu đã thu được Recall đáng kể; MostPopular làm trực tiếp việc đó, còn MF có hệ số chệch riêng cho từng item.")
P(f"**RippleNet xếp hạng kém nhất ở các K nhỏ trên cả ba tập.** Recall@10 của RippleNet là {f4(mv['RippleNet']['recall']['10'])} ở tập phim, {f4(bk['RippleNet']['recall']['10'])} ở tập sách và {f4(ms['RippleNet']['recall']['10'])} ở tập nhạc, dù AUC của nó không thua xa các mô hình khác. Ở K lớn khoảng cách thu hẹp: Recall@100 trên tập phim là {f4(mv['RippleNet']['recall']['100'])}, và trên tập sách RippleNet vượt CKAN từ K = 20. Đề tài chưa xác định được nguyên nhân bằng thực nghiệm. Một giả thuyết phù hợp với kiến trúc là: item trong RippleNet chỉ là một vector tra bảng với số chiều nhỏ, nên khi xếp hạng trên toàn danh mục, một số item hiếm có embedding ít được huấn luyện vẫn nhận điểm cao và chiếm các vị trí đầu.")

H2("Kết quả theo độ thưa dữ liệu")
t, f = R.tab(), R.fig()
P(f"{t} và {f} trình bày AUC khi chỉ giữ lại một phần tập huấn luyện. AUC được đo trên nhóm người dùng đã có tương tác dương ở mức 10%: {vn(DS['movie']['sparsity']['evalUsers'])} người dùng ở tập phim, {vn(DS['book']['sparsity']['evalUsers'])} ở tập sách và {vn(DS['music']['sparsity']['evalUsers'])} ở tập nhạc. Vì vậy cột 100% của bảng này không trùng với Bảng 6.2, vốn đo trên toàn bộ tập kiểm thử.", keep=True)
ratios = DS["movie"]["sparsity"]["ratios"]
rows, bold = [], set()
for d in DS:
    sp = DS[d]["sparsity"]["auc"]
    for i in range(len(ratios)):
        best = max(sp[m_][i] for m_ in ("MF", "RippleNet", "CKAN"))
        for m_i, m_ in enumerate(("MF", "RippleNet", "CKAN")):
            if sp[m_][i] == best:
                bold.add((len(rows) + m_i, 2 + i))
    for m_ in ("MF", "RippleNet", "CKAN"):
        rows.append([NAMES[d], m_] + [f4(v) for v in sp[m_]])
R.table(t, "ROC-AUC trên tập kiểm thử theo tỷ lệ tập huấn luyện được giữ lại", ["Tập dữ liệu", "Mô hình"] + [f"{int(r * 100)}%" for r in ratios], rows, [3.0, 2.3, 1.75, 1.75, 1.75, 1.75, 1.75, 1.75], left_cols=(0, 1), bold_cells=bold, size=10)
fig("chart_sparsity", f, "ROC-AUC theo tỷ lệ tập huấn luyện được giữ lại", 16.0)
f = R.fig()
sp_mv, sp_bk, sp_ms = (DS[d]["sparsity"]["auc"] for d in ("movie", "book", "music"))
P(f"**Kết quả khác nhau rõ giữa ba tập; đồ thị tri thức không tự động bù được việc thiếu tương tác.** {f} cho thấy chênh lệch giữa CKAN và MF ở từng mức dữ liệu.")
fig("chart_gain", f, "Chênh lệch tương đối về AUC giữa CKAN và MF theo lượng dữ liệu huấn luyện", 11.5)
P(f"**Tập phim: hai mô hình đồ thị chịu thiếu dữ liệu tốt hơn MF.** Với 10% dữ liệu, CKAN đạt {f4(sp_mv['CKAN'][0])} và RippleNet đạt {f4(sp_mv['RippleNet'][0])}, so với {f4(sp_mv['MF'][0])} của MF; CKAN cao hơn MF {signed(gain('movie', 0))}. Thứ tự CKAN, RippleNet, MF giữ nguyên đến mức 80%, và khoảng cách thu hẹp dần cho đến khi ba mô hình ngang nhau ở 100% (chênh lệch {signed(gain('movie', 5))}). Đây là tập có đồ thị tri thức dày nhất, mỗi phim có trung vị {vn(m[0])} bộ ba.")
P(f"**Tập sách: CKAN dẫn đầu ở mọi mức, nhưng ở 10% gần như ngang MF.** Với 10% dữ liệu, CKAN đạt {f4(sp_bk['CKAN'][0])} và MF đạt {f4(sp_bk['MF'][0])}, còn RippleNet chỉ {f4(sp_bk['RippleNet'][0])}, tức ngang mức đoán ngẫu nhiên. Từ 20% trở đi CKAN tách khỏi MF, và khoảng cách còn {signed(gain('book', 5))} khi có đủ dữ liệu.")
P(f"**Tập nhạc: MF tốt hơn cả hai mô hình đồ thị khi dữ liệu ít.** Ở mức 10%, MF đạt {f4(sp_ms['MF'][0])} trong khi CKAN chỉ {f4(sp_ms['CKAN'][0])} và RippleNet {f4(sp_ms['RippleNet'][0])}; ở mức 20% MF vẫn dẫn đầu. CKAN vượt MF từ mức 40% và hơn {signed(gain('music', 5))} khi có đủ dữ liệu. RippleNet thấp hơn MF ở mọi mức.")
P(f"**Giải thích.** Hai mô hình đồ thị biểu diễn người dùng hoàn toàn bằng các item họ đã tương tác và các bộ ba xuất phát từ đó. Khi người dùng chỉ còn một hai lượt tương tác và mỗi item chỉ có trung vị {vn(m[1])} bộ ba như ở tập sách và tập nhạc, tập bộ ba của người dùng gần như không có nội dung, và mô hình cũng không đủ mẫu để học embedding cho các thực thể. MF thì có hệ số chệch riêng cho từng item, học được độ phổ biến của item ngay cả khi dữ liệu ít; với cách lấy mẫu âm đều, chỉ riêng độ phổ biến đã cho AUC khá cao (xem kết quả của MostPopular ở Bảng 6.2). Ở tập phim, đồ thị dày hơn nhiều nên vài lượt tương tác còn lại vẫn dẫn tới đủ thực thể để mô tả người dùng. Đây là cách giải thích phù hợp với số liệu; đề tài chưa kiểm chứng riêng từng yếu tố.")
P(f"**CKAN so với RippleNet.** CKAN cao hơn RippleNet ở mọi mức dữ liệu trên cả ba tập, trừ mức 100% của tập phim nơi hai mô hình ngang nhau ({f4(sp_mv['CKAN'][5])} và {f4(sp_mv['RippleNet'][5])}). Khoảng cách lớn nhất nằm ở các mức thấp của tập sách và tập nhạc, nơi RippleNet gần mức đoán ngẫu nhiên.")

H2("Thảo luận và giới hạn của thực nghiệm")
B([
    "**Lợi ích của đồ thị tri thức phụ thuộc vào độ dày của chính đồ thị.** Ở tập phim, hai mô hình đồ thị hơn MF rõ nhất khi tương tác ít. Ở tập sách và tập nhạc, nơi mỗi item chỉ có một vài bộ ba, lợi thế đó không xuất hiện ở các mức dữ liệu thấp nhất.",
    "**CKAN hơn RippleNet rõ ở tập sách và tập nhạc**, về AUC, về Recall ở các K nhỏ và ở mọi mức dữ liệu của thí nghiệm độ thưa; ở tập phim hai mô hình ngang nhau về AUC. Lan truyền hai phía cùng tín hiệu cộng tác là phần tạo ra khác biệt, vì RippleNet tự nó không vượt được MF ở hai tập thưa.",
    "**AUC và Recall@K đo hai việc khác nhau.** CKAN có AUC cao nhất ở tập sách nhưng Recall thấp hơn cả MostPopular; kết luận về mô hình cần nêu rõ dựa trên độ đo nào.",
    "**So với bài báo.** Số liệu của đề tài không so trực tiếp được với bảng kết quả trong hai bài báo: tập phim chỉ dùng 2.500 người dùng, RippleNet trên tập nhạc dùng cấu hình mượn từ tập phim, và mỗi mô hình chỉ huấn luyện tối đa 10 epoch.",
    "**Mỗi cấu hình chỉ chạy một lần** với một hạt giống ngẫu nhiên, nên các chênh lệch nhỏ (như giữa bốn mô hình trên tập phim) chưa đủ cơ sở để kết luận hơn kém.",
    "**Recall@K được đo trên 100 người dùng**, đủ để so sánh xu hướng nhưng có độ dao động đáng kể.",
    "**Cách lấy mẫu âm đều làm AUC ưu ái độ phổ biến**, như đã phân tích với MostPopular.",
])

# =============================================================================== chapter 7
H1("XÂY DỰNG HỆ THỐNG MINH HOẠ")
P("Để kiểm chứng mô hình trong điều kiện sử dụng và trình bày trực quan các bài toán của đề tài, nhóm xây dựng một ứng dụng web dùng mô hình CKAN do notebook thực nghiệm xuất ra.")
H2("Kiến trúc hệ thống")
f, t = R.fig(), R.tab()
P(f"Hệ thống gồm một giao diện web, một dịch vụ API và các kho dữ liệu, như {f}. {t} liệt kê công nghệ của từng thành phần.")
fig("diagram_system", f, "Kiến trúc hệ thống minh hoạ", 15.0)
R.table(t, "Các thành phần của hệ thống", ["Thành phần", "Công nghệ", "Vai trò"], [
    ["Giao diện", "React 18, TypeScript, Vite, Tailwind CSS", "Các màn hình gợi ý Top-K, đồ thị tri thức, kết quả thực nghiệm và thư viện"],
    ["Dịch vụ API", "FastAPI, Pydantic", "Nhận yêu cầu từ giao diện, kiểm tra tham số đầu vào"],
    ["Bộ máy gợi ý", "PyTorch, NumPy", "Nạp mô hình CKAN của ba tập, chấm điểm, tìm đường dẫn lý giải"],
    ["Dữ liệu người dùng (tập phim)", "PostgreSQL", "Người dùng, đánh giá, tương tác"],
    ["Đồ thị (tập phim)", "Neo4j", "Truy vấn đồ thị con quanh người dùng"],
    ["Mô hình và dữ liệu", "Tệp do notebook xuất", "Trọng số, cấu hình, ma trận embedding item, đồ thị tri thức, metadata"],
], [3.8, 4.6, 7.6], left_cols=(0, 1, 2))

H2("Luồng suy luận")
f = R.fig()
P(f"Như đã phân tích ở Chương 4, biểu diễn item của CKAN không phụ thuộc người dùng. Hệ thống tận dụng tính chất này ({f}): ma trận embedding của toàn bộ item được notebook tính sẵn một lần; mỗi yêu cầu gợi ý chỉ cần dựng tập bộ ba của người dùng từ các item họ đã thích, đưa qua lớp attention để có biểu diễn người dùng, rồi nhân với ma trận item và lấy sigmoid.")
fig("diagram_inference", f, "Luồng suy luận: phần item tính trước, mỗi yêu cầu chỉ tính phần người dùng", 15.5)
B([
    "**Chấm điểm toàn bộ danh mục.** Nhờ ma trận tính trước, mỗi yêu cầu xếp hạng trên toàn bộ danh mục (gần 17.000 phim) thay vì một tập ứng viên rút gọn; phần chấm điểm của bộ máy mất dưới 10 mili giây trên CPU trong các lần đo của nhóm.",
    "**Khớp với mô hình thực nghiệm.** Hệ thống nạp đúng trọng số, số tầng lan truyền và kích thước tập bộ ba của từng tập dữ liệu như trong notebook. Nhóm đã kiểm tra bằng cách tính lại embedding của một số item từ tập bộ ba do notebook lưu và so với ma trận notebook xuất ra: sai lệch lớn nhất ở mức 10⁻⁷ trên cả ba tập.",
    "**Cập nhật tức thời.** Khi người dùng thích hoặc ẩn một item trên giao diện, lần gợi ý kế tiếp dùng ngay tập hạt giống mới mà không cần huấn luyện lại, vì người dùng được biểu diễn bằng các item đã tương tác chứ không bằng một vector riêng.",
    "**Người dùng chưa có tương tác** nhận danh sách phổ biến nhất của tập dữ liệu; hệ thống nói rõ đây không phải gợi ý cá nhân hoá.",
])
P("Có hai khác biệt so với thực nghiệm cần nêu. Thứ nhất, mô hình đang chạy trong hệ thống là bản CKAN do notebook xuất ra ở lần huấn luyện trước khi bổ sung lan truyền cộng tác phía item, tức tập khởi đầu của mỗi item chỉ gồm chính nó; các số đo ở Chương 6 thuộc bản đã bổ sung. Thứ hai, 2.500 người dùng tập phim trong cơ sở dữ liệu của hệ thống không trùng với 2.500 người dùng được lấy mẫu trong notebook. Điều này không ảnh hưởng đến tính đúng của mô hình vì CKAN không có tham số riêng cho từng người dùng, nhưng có nghĩa là các người dùng trong phần minh hoạ không phải những người dùng đã được dùng để huấn luyện và đánh giá. Người dùng của tập sách và tập nhạc thì trùng nhau.")

H2("Lý giải gợi ý")
P("Với mỗi gợi ý, hệ thống tìm các đường dẫn hai bước trên đồ thị tri thức có dạng: item đã thích → thực thể ← item được gợi ý. Mỗi đường dẫn được trình bày bằng lời (ví dụ “cùng tác giả với …”) và trên một đồ thị con tương tác. Ba nguyên tắc được tuân thủ để lời lý giải trung thực:")
B([
    "**Điểm số là điểm của mô hình.** Điểm hiển thị trong bảng lý giải được CKAN tính cho đúng cặp người dùng và item, không phải một hằng số.",
    "**Câu hỏi “nếu bỏ lượt thích này” được trả lời bằng phép đo.** Hệ thống chấm lại điểm sau khi loại item nguồn của đường dẫn đầu tiên khỏi lịch sử và báo cả hai giá trị.",
    f"**Không có đường dẫn thì nói là không có.** Như Chương 5 đã chỉ ra, chỉ khoảng {vn(cover_[1], 0)}% gợi ý ở tập sách và {vn(cover_[2], 0)}% ở tập nhạc có đường dẫn qua thực thể chung. Với các gợi ý còn lại, hệ thống thông báo rằng điểm số đến từ embedding mà mô hình học được từ hành vi người dùng, thay vì dựng một lời giải thích không có thật.",
])
P("Đồ thị tri thức gốc chỉ chứa mã số thực thể. Để đường dẫn đọc được, nhóm khôi phục tên thực thể bằng cách đối chiếu: nối mã phim MovieLens sang IMDb, lấy danh sách đạo diễn, diễn viên, biên kịch của từng phim từ Wikidata, rồi đặt tên cho một thực thể khi đa số các phim gắn với nó cùng ghi một tên ở đúng vai trò đó. Cách này đặt được tên cho 33.462 trên 85.615 thực thể của đồ thị phim, phủ 59% số bộ ba. Với tập sách, tên tác giả lấy từ metadata có sẵn, đặt được 7.229 thực thể. Các thực thể còn lại được hiển thị theo vai trò kèm mã số, ví dụ “Thể loại #17137”. Ảnh nghệ sĩ của tập nhạc được lấy từ dịch vụ Deezer (3.711 trên 3.846 nghệ sĩ) vì các đường dẫn ảnh trong bộ dữ liệu gốc không còn truy cập được.")

H2("Giao diện và các kịch bản sử dụng")
P("Thanh trên cùng của giao diện cho phép đổi tập dữ liệu và đổi người dùng ở bất kỳ màn hình nào; màu nhấn đổi theo tập dữ liệu đang chọn. Phần này đi qua các màn hình chính: gợi ý Top-K, bảng lý giải, đồ thị tri thức, thư viện và chọn người dùng. Các hình là ảnh chụp trực tiếp từ hệ thống đang chạy với mô hình và dữ liệu thật.")

H3("Màn hình gợi ý Top-K")
f = R.fig()
P(f"Màn hình đầu tiên trả lời bài toán xếp hạng Top-K. Phần đầu trang nêu quy mô tập dữ liệu; tiếp theo là gợi ý phù hợp nhất kèm lý do; bên dưới là các gợi ý xếp hạng 2 đến 15, mỗi thẻ có thứ hạng, điểm dự đoán của CKAN, lý do ngắn và hai nút thích, ẩn ({f}; ảnh chụp phần đầu trang).")
shot("rec_movie", f, "Gợi ý phim cho người dùng #1 của tập MovieLens-20M", WIDE, crop=0.66)
f = R.fig()
P(f"Đổi tập dữ liệu trên thanh điều hướng sẽ chuyển sang mô hình CKAN của tập đó và người dùng mẫu tương ứng. {f} là gợi ý sách: lý do đi kèm mỗi cuốn được lấy từ đường dẫn tri thức, ví dụ cùng tác giả hoặc cùng thể loại với một cuốn đã thích.")
shot("rec_book", f, "Gợi ý sách cho người dùng #790 của tập Book-Crossing", WIDE, crop=0.66)
f = R.fig()
P(f"{f} là gợi ý nghệ sĩ ở tập nhạc. Nhiều thẻ ghi lý do chung chung “phù hợp với đặc trưng sở thích”, phản ánh việc đồ thị tri thức của tập này mỏng và phần lớn gợi ý không có đường dẫn tri thức (Mục 5.4).")
shot("rec_music", f, "Gợi ý nghệ sĩ cho người dùng #774 của tập Last.FM", WIDE, crop=0.66)

H3("Bảng lý giải một gợi ý")
f = R.fig()
P(f"Bấm vào một gợi ý sẽ mở bảng lý giải ở cạnh phải ({f}). Bảng gồm câu tóm tắt, độ tin cậy tính theo số đường dẫn tìm được, điểm CKAN của đúng cặp người dùng và item, đồ thị con tương tác và danh sách các đường dẫn tri thức.")
shot("explain_movie", f, "Bảng lý giải cho gợi ý đứng đầu của người dùng #1: đồ thị con và các đường dẫn tri thức", WIDE)
f = R.fig()
P(f"Bảng có thể mở rộng toàn màn hình để xem song song đồ thị con, các đường dẫn, tỷ lệ đường dẫn theo loại quan hệ và kết quả phép thử bỏ lượt thích ({f}). Tên các thực thể như biên kịch hay hãng sản xuất là tên thật đã được khôi phục.")
shot("explain_movie_wide", f, "Bảng lý giải ở chế độ mở rộng: đồ thị con, đường dẫn, tỷ lệ theo loại quan hệ", WIDE)
f = R.fig()
P(f"{f} là một lý giải ở tập sách, nơi các đường dẫn chủ yếu đi qua quan hệ tác giả.")
shot("explain_book", f, "Lý giải một gợi ý sách qua quan hệ tác giả", WIDE)
f = R.fig()
P(f"{f} là trường hợp không có đường dẫn ở tập nhạc: hệ thống báo rõ điều đó, xếp độ tin cậy ở mức thấp và cho biết điểm số đến từ embedding của mô hình.")
shot("explain_music_no_path", f, "Trường hợp không có đường dẫn tri thức: hệ thống nói rõ thay vì dựng lời giải thích", WIDE)

H3("Màn hình đồ thị tri thức")
f = R.fig()
P(f"Màn hình thứ hai trả lời bài toán đường dẫn tri thức ở mức toàn cảnh: người dùng ở trung tâm, nối tới các item đã thích, từ đó qua các thực thể chung tới các item được gợi ý ({f}). Đồ thị được bố trí bằng mô phỏng lực; rê chuột vào một nút sẽ chỉ giữ lại các liên kết của nút đó, bấm vào nút để xem chi tiết.")
shot("graph_movie", f, "Đồ thị tri thức quanh người dùng #1 của tập phim", WIDE, crop=0.72)
f = R.fig()
P(f"Bộ lọc theo loại thực thể được sinh từ chính dữ liệu đang hiển thị. {f} là đồ thị sau khi chỉ giữ các thực thể đạo diễn.")
shot("graph_movie_director", f, "Đồ thị sau khi lọc chỉ giữ thực thể đạo diễn", WIDE, crop=0.72)
f = R.fig()
P(f"Với tập sách và tập nhạc, đồ thị được dựng từ kết quả lý giải của gợi ý đứng đầu ({f}).")
shot("graph_book", f, "Đồ thị quanh người dùng #790 của tập sách", WIDE, crop=0.72)

H3("Thư viện")
f = R.fig()
P(f"Màn hình thư viện cho phép duyệt toàn bộ danh mục và thích các item để bổ sung vào hồ sơ của người dùng đang chọn ({f}). Các lượt thích được lưu lại và có hiệu lực ngay ở lần gợi ý kế tiếp.")
shot("catalog_movie", f, "Thư viện phim", WIDE)
f = R.fig()
P(f"Ô tìm kiếm lọc theo tên item hoặc tác giả. {f} là kết quả tìm “tolkien” ở tập sách, với một cuốn vừa được thích.")
shot("catalog_book_search", f, "Tìm kiếm trong thư viện sách và thích một cuốn", WIDE)

H3("Chọn và tạo người dùng")
f = R.fig()
P(f"Hộp thoại chọn người dùng liệt kê các hồ sơ mẫu của tập dữ liệu, cho phép tìm theo mã số hoặc duyệt toàn bộ ({f}).")
shot("users_movie", f, "Hộp thoại chọn người dùng của tập phim", 13.5)
f = R.fig()
P(f"Người xem cũng có thể tạo một người dùng mới hoàn toàn chưa có tương tác ({f}); người dùng này nhận mã số kế tiếp của đúng tập dữ liệu đang chọn.")
shot("users_music_new", f, "Tạo người dùng mới ở tập nhạc", 13.5)
f = R.fig()
P(f"Vì chưa có lượt thích nào, người dùng mới nhận danh sách phổ biến nhất ({f}); chỉ cần thích một item trong thư viện là CKAN bắt đầu gợi ý riêng. Đây là kịch bản khởi động lạnh có thể trình diễn trực tiếp.")
shot("rec_music_new_user", f, "Người dùng mới chưa có lượt thích nhận danh sách phổ biến nhất", WIDE, crop=0.66)

# =============================================================================== chapter 8
H1("KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN")
H2("Kết quả đạt được")
B([
    "Đã nghiên cứu và trình bày chi tiết hai mô hình RippleNet và CKAN, làm rõ ba khác biệt cốt lõi: lan truyền hai phía, mã hoá tín hiệu cộng tác và attention nhận biết tri thức.",
    "Đã cài đặt lại hai mô hình cùng hai phương pháp cơ sở và thực nghiệm trên ba tập dữ liệu thuộc ba miền với ba nhóm đánh giá.",
    f"Kết quả cho thấy CKAN dẫn đầu về dự đoán CTR trên tập sách (AUC {f4(bk['CKAN']['auc'])}) và tập nhạc (AUC {f4(ms['CKAN']['auc'])}), và dẫn đầu về xếp hạng Top-K trên tập nhạc từ K = 10; trên tập phim, nơi dữ liệu tương tác dày, các mô hình gần như ngang nhau về AUC. Về Recall@K, MF dẫn đầu ở tập phim và MostPopular ở tập sách.",
    f"Thí nghiệm độ thưa cho kết quả khác nhau theo tập dữ liệu. Với 10% dữ liệu huấn luyện, CKAN cao hơn MF {signed(gain('movie', 0))} về AUC ở tập phim, ngang MF ở tập sách ({signed(gain('book', 0))}) và thấp hơn MF ở tập nhạc ({signed(gain('music', 0))}); CKAN chỉ vượt MF ở tập nhạc từ mức 40%. Đồ thị tri thức giúp chịu thiếu dữ liệu khi bản thân đồ thị đủ dày.",
    "Đã xây dựng hệ thống web minh hoạ dùng đúng mô hình thực nghiệm, có lý giải từng gợi ý bằng đường dẫn tri thức thật và trình bày lại kết quả thực nghiệm.",
])
H2("Hạn chế")
B([
    "Tập phim chỉ dùng 2.500 người dùng và RippleNet trên tập nhạc dùng cấu hình mượn từ tập phim, nên số liệu không so trực tiếp được với hai bài báo.",
    "Mỗi cấu hình chỉ chạy một lần; Recall@K đo trên 100 người dùng. Nguyên nhân RippleNet xếp hạng kém ở các K nhỏ và nguyên nhân hai mô hình đồ thị thua MF ở các mức dữ liệu thấp của tập nhạc mới dừng ở giả thuyết.",
    "Mô hình trong hệ thống minh hoạ là bản CKAN của lần huấn luyện trước, chưa có lan truyền cộng tác phía item.",
    "Đồ thị tri thức của tập nhạc và tập sách mỏng, nên phần lớn gợi ý ở hai tập này không có đường dẫn tri thức để lý giải.",
    "Tên thực thể mới khôi phục được một phần; người dùng tập phim trong hệ thống không trùng với mẫu người dùng của thực nghiệm.",
])
H2("Hướng phát triển")
B([
    "Lặp lại thực nghiệm với nhiều hạt giống và báo cáo độ lệch chuẩn; tách riêng đóng góp của lan truyền cộng tác phía item bằng cách tắt thành phần này và đo lại.",
    "Kiểm chứng các giả thuyết ở Chương 6: thêm hệ số chệch theo item vào CKAN để xem có cải thiện Recall@K và kết quả ở các mức dữ liệu thấp hay không; tăng số chiều của RippleNet để xem ảnh hưởng tới xếp hạng.",
    "Đưa mô hình CKAN của lần huấn luyện mới vào hệ thống minh hoạ.",
    "Mở rộng tìm đường dẫn lý giải tới hai bước ở mỗi phía cho khớp số tầng của mô hình, và bổ sung lý giải dựa trên người dùng chung cho các gợi ý không có đường dẫn tri thức.",
    "Thử nghiệm các mô hình kế tiếp như KGAT hoặc KGIN trên cùng quy trình để mở rộng so sánh.",
])

# =============================================================================== references
H1("TÀI LIỆU THAM KHẢO", numbered=False)
REFS = [
    'H. Wang, F. Zhang, J. Wang, M. Zhao, W. Li, X. Xie, and M. Guo, "RippleNet: Propagating User Preferences on the Knowledge Graph for Recommender Systems," in *Proceedings of the 27th ACM International Conference on Information and Knowledge Management (CIKM)*, 2018, pp. 417–426.',
    'Z. Wang, G. Lin, H. Tan, Q. Chen, and X. Liu, "CKAN: Collaborative Knowledge-aware Attentive Network for Recommender Systems," in *Proceedings of the 43rd International ACM SIGIR Conference on Research and Development in Information Retrieval (SIGIR)*, 2020, pp. 219–228.',
    'H. Wang, M. Zhao, X. Xie, W. Li, and M. Guo, "Knowledge Graph Convolutional Networks for Recommender Systems," in *Proceedings of The World Wide Web Conference (WWW)*, 2019.',
    'X. Wang, X. He, Y. Cao, M. Liu, and T.-S. Chua, "KGAT: Knowledge Graph Attention Network for Recommendation," in *Proceedings of the 25th ACM SIGKDD International Conference on Knowledge Discovery and Data Mining (KDD)*, 2019, pp. 950–958.',
    'F. Zhang, N. J. Yuan, D. Lian, X. Xie, and W.-Y. Ma, "Collaborative Knowledge Base Embedding for Recommender Systems," in *Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining (KDD)*, 2016, pp. 353–362.',
    'H. Wang, F. Zhang, X. Xie, and M. Guo, "DKN: Deep Knowledge-Aware Network for News Recommendation," in *Proceedings of the 2018 World Wide Web Conference (WWW)*, 2018, pp. 1835–1844.',
    'Y. Koren, R. Bell, and C. Volinsky, "Matrix Factorization Techniques for Recommender Systems," *Computer*, vol. 42, no. 8, pp. 30–37, 2009.',
    'A. Bordes, N. Usunier, A. García-Durán, J. Weston, and O. Yakhnenko, "Translating Embeddings for Modeling Multi-relational Data," in *Advances in Neural Information Processing Systems (NIPS)*, 2013.',
    'X. Yu, X. Ren, Y. Sun, Q. Gu, B. Sturt, U. Khandelwal, B. Norick, and J. Han, "Personalized Entity Recommendation: A Heterogeneous Information Network Approach," in *Proceedings of the 7th ACM International Conference on Web Search and Data Mining (WSDM)*, 2014.',
    'Q. Guo, F. Zhuang, C. Qin, H. Zhu, X. Xie, H. Xiong, and Q. He, "A Survey on Knowledge Graph-Based Recommender Systems," *IEEE Transactions on Knowledge and Data Engineering*, vol. 34, no. 8, 2022.',
    'F. M. Harper and J. A. Konstan, "The MovieLens Datasets: History and Context," *ACM Transactions on Interactive Intelligent Systems*, vol. 5, no. 4, 2015.',
    'C.-N. Ziegler, S. M. McNee, J. A. Konstan, and G. Lausen, "Improving Recommendation Lists Through Topic Diversification," in *Proceedings of the 14th International Conference on World Wide Web (WWW)*, 2005.',
    'I. Cantador, P. Brusilovsky, and T. Kuflik, "Second Workshop on Information Heterogeneity and Fusion in Recommender Systems (HetRec 2011)," in *Proceedings of the 5th ACM Conference on Recommender Systems (RecSys)*, 2011.',
    'D. P. Kingma and J. Ba, "Adam: A Method for Stochastic Optimization," in *International Conference on Learning Representations (ICLR)*, 2015.',
    'Mã nguồn RippleNet do tác giả công bố: https://github.com/hwwang55/RippleNet (truy cập tháng 10/2026).',
    'Mã nguồn CKAN do tác giả công bố: https://github.com/weberrr/CKAN (truy cập tháng 10/2026).',
]
for i, ref in enumerate(REFS, 1):
    para = P(f"[{i}] {ref}", indent=False, size=12)
    para.paragraph_format.left_indent = Cm(0.9)
    para.paragraph_format.first_line_indent = Cm(-0.9)

R.save(OUT)
print("saved", OUT, f"{OUT.stat().st_size / 1e6:.1f} MB")
