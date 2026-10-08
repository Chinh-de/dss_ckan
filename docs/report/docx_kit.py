"""Small python-docx layer for the report: page setup, styles, captions, tables, figures, equations.

Formatting follows the report standard used for this project: A4, margins 3/2/2/2 cm,
Times New Roman throughout, justified body text at 1.3 line spacing, navy table headers
with zebra rows, captions above tables and below figures.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

FONT = "Times New Roman"
BODY_PT = 13
NAVY, BLUE, SLATE, GRAY = "0F172A", "1E3A8A", "334155", "64748B"
TEXT_WIDTH_CM = 16.0
FIG_STYLE, TAB_STYLE = "Chú thích hình", "Chú thích bảng"


def _set_font(style_or_run, size=None, bold=None, italic=None, color=None):
    font = style_or_run.font
    font.name = FONT
    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    if color is not None:
        font.color.rgb = RGBColor.from_string(color)


def _field(paragraph, instruction, placeholder=""):
    """Insert a Word field (TOC, PAGE...) that Word fills in when fields are updated."""
    def fld(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        return el

    run = paragraph.add_run()
    run._r.append(fld("begin"))
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run._r.append(instr)
    run._r.append(fld("separate"))
    if placeholder:
        text = OxmlElement("w:t")
        text.text = placeholder
        run._r.append(text)
    run._r.append(fld("end"))
    return run


class Counter:
    """Hands out 'Hình 3.2' style labels, numbered within the current chapter."""

    def __init__(self, word, report):
        self.word, self.report, self.n = word, report, 0

    def __call__(self):
        self.n += 1
        return f"{self.word} {self.report.chapter}.{self.n}"


class Report:
    def __init__(self):
        self.doc = Document()
        self.chapter = 0
        self.fig = Counter("Hình", self)
        self.tab = Counter("Bảng", self)
        self._eq = 0
        self._page_setup()
        self._styles()

    # ------------------------------------------------------------------ setup
    def _page_setup(self):
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.left_margin, section.right_margin = Cm(3.0), Cm(2.0)
        section.top_margin, section.bottom_margin = Cm(2.0), Cm(2.0)
        # Ask Word to refresh the table of contents and figure/table lists on open.
        settings = self.doc.settings.element
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)

    def _styles(self):
        styles = self.doc.styles
        normal = styles["Normal"]
        _set_font(normal, size=BODY_PT)
        pf = normal.paragraph_format
        pf.line_spacing, pf.space_before, pf.space_after = 1.3, Pt(0), Pt(6)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        for name, size, color, before, after in [
            ("Heading 1", 16, NAVY, 18, 8),
            ("Heading 2", 14, BLUE, 14, 6),
            ("Heading 3", 12, SLATE, 10, 4),
        ]:
            st = styles[name]
            _set_font(st, size=size, bold=True, italic=False, color=color)
            pf = st.paragraph_format
            pf.space_before, pf.space_after = Pt(before), Pt(after)
            pf.keep_with_next = True
            pf.line_spacing = 1.2
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        styles["Heading 1"].paragraph_format.page_break_before = True

        for name, bold in [(FIG_STYLE, False), (TAB_STYLE, True)]:
            st = styles.add_style(name, 1)
            st.base_style = normal
            _set_font(st, size=10.5, italic=True, bold=bold, color=SLATE if bold else GRAY)
            pf = st.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.line_spacing = 1.15
            pf.space_before, pf.space_after = (Pt(10), Pt(4)) if bold else (Pt(4), Pt(12))
            pf.keep_with_next = bold

        for name in ("List Bullet", "List Number"):
            _set_font(styles[name], size=BODY_PT)
            styles[name].paragraph_format.space_after = Pt(3)
            styles[name].paragraph_format.line_spacing = 1.3

        for name in ("TOC Heading", "toc 1", "toc 2", "toc 3", "table of figures"):
            try:
                _set_font(styles[name], size=12)
            except KeyError:
                pass

    def page_numbers(self):
        """Centered page number in the footer of the current (last) section."""
        footer = self.doc.sections[-1].footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(_field(p, "PAGE", "1"), size=11)

    def new_section(self, restart_numbering=False):
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        if restart_numbering:
            pg = OxmlElement("w:pgNumType")
            pg.set(qn("w:start"), "1")
            section._sectPr.append(pg)
        return section

    # ------------------------------------------------------------------ text
    INLINE = re.compile(r"(\*\*.+?\*\*|\$.+?\$|\*.+?\*)")

    def _math_runs(self, paragraph, expr, size, bold=None):
        """Inline maths: letters italic, _{x} subscript, ^{x} superscript."""
        i = 0
        while i < len(expr):
            ch = expr[i]
            if ch in "_^" and i + 1 < len(expr):
                if expr[i + 1] == "{":
                    end = expr.index("}", i)
                    token, nxt = expr[i + 2:end], end + 1
                else:
                    token, nxt = expr[i + 1], i + 2
                run = paragraph.add_run(token)
                _set_font(run, size=size, italic=True, bold=bold)
                run.font.subscript = ch == "_"
                run.font.superscript = ch == "^"
                i = nxt
                continue
            j = i
            while j < len(expr) and expr[j] not in "_^":
                j += 1
            run = paragraph.add_run(expr[i:j])
            _set_font(run, size=size, italic=True, bold=bold)
            i = j

    def runs(self, paragraph, text, size=BODY_PT, bold=None, color=None):
        for part in self.INLINE.split(text):
            if not part:
                continue
            if part.startswith("**"):
                # Bold spans may themselves contain inline maths.
                self.runs(paragraph, part[2:-2], size=size, bold=True, color=color)
            elif part.startswith("$"):
                self._math_runs(paragraph, part[1:-1], size, bold)
            elif part.startswith("*"):
                _set_font(paragraph.add_run(part[1:-1]), size=size, italic=True, bold=bold, color=color)
            else:
                _set_font(paragraph.add_run(part), size=size, bold=bold, color=color)

    def p(self, text, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=BODY_PT, bold=None, color=None, after=None, keep=False):
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = align
        if indent and align == WD_ALIGN_PARAGRAPH.JUSTIFY:
            paragraph.paragraph_format.first_line_indent = Cm(1.0)
        if after is not None:
            paragraph.paragraph_format.space_after = Pt(after)
        paragraph.paragraph_format.keep_with_next = keep
        self.runs(paragraph, text, size=size, bold=bold, color=color)
        return paragraph

    def bullets(self, items, numbered=False):
        for item in items:
            paragraph = self.doc.add_paragraph(style="List Number" if numbered else "List Bullet")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            self.runs(paragraph, item)

    def h1(self, title, numbered=True):
        if numbered:
            self.chapter += 1
            self.fig.n = self.tab.n = 0
            self._h2 = self._h3 = 0
            title = f"CHƯƠNG {self.chapter}. {title}"
        heading = self.doc.add_heading(level=1)
        _set_font(heading.add_run(title), size=16, bold=True, color=NAVY)
        return heading

    def h2(self, title):
        self._h2 += 1
        self._h3 = 0
        heading = self.doc.add_heading(level=2)
        _set_font(heading.add_run(f"{self.chapter}.{self._h2}. {title}"), size=14, bold=True, color=BLUE)

    def h3(self, title):
        self._h3 += 1
        heading = self.doc.add_heading(level=3)
        _set_font(heading.add_run(f"{self.chapter}.{self._h2}.{self._h3}. {title}"), size=12, bold=True, color=SLATE)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def toc(self, instruction):
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _field(paragraph, instruction, "Nhấn F9 trong Word để cập nhật danh mục.")

    # ------------------------------------------------------------------ figures
    def figure(self, path, label, caption, width_cm=15.5):
        path = Path(path)
        with Image.open(path) as im:
            ratio = im.height / im.width
        # Keep tall images on one page together with their caption.
        width_cm = min(width_cm, TEXT_WIDTH_CM, 20.5 / ratio)
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
        cap = self.doc.add_paragraph(style=FIG_STYLE)
        self.runs(cap, f"{label}: {caption}", size=10.5, color=GRAY)
        for run in cap.runs:
            run.font.italic = True

    def equation_label(self):
        """Reserve the next equation number, e.g. '(3.2)', so the text can cite it first."""
        self._eq = self._eq + 1 if getattr(self, "_eq_chapter", None) == self.chapter else 1
        self._eq_chapter = self.chapter
        return f"({self.chapter}.{self._eq})"

    def equation(self, path, label, height_scale=1.0):
        """Centered display equation with its number right-aligned on the same line."""
        with Image.open(path) as im:
            width_cm = im.width / 300 * 2.54 * height_scale
        width_cm = min(width_cm, TEXT_WIDTH_CM - 1.6)
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = paragraph.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(4), Pt(8), 1.0
        pf.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM / 2), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
        paragraph.add_run("\t").add_picture(str(path), width=Cm(width_cm))
        _set_font(paragraph.add_run("\t" + label), size=BODY_PT)
        return label

    # ------------------------------------------------------------------ tables
    @staticmethod
    def _shade(cell, fill):
        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tcpr.append(shd)

    @staticmethod
    def _margins(cell, dxa=100):
        tcpr = cell._tc.get_or_add_tcPr()
        mar = OxmlElement("w:tcMar")
        for side in ("top", "bottom", "start", "end"):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:w"), str(dxa if side in ("start", "end") else 80))
            el.set(qn("w:type"), "dxa")
            mar.append(el)
        tcpr.append(mar)

    @staticmethod
    def _borders(table, color="CBD5E1"):
        tblpr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
            borders.append(el)
        tblpr.append(borders)

    def table(self, label, caption, header, rows, widths_cm, left_cols=(0,), bold_cells=(), size=10.5):
        """Caption above, navy header, zebra rows. `bold_cells` is a set of (row, col) to emphasise."""
        cap = self.doc.add_paragraph(style=TAB_STYLE)
        self.runs(cap, f"{label}: {caption}", size=10.5, bold=True, color=SLATE)

        table = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        self._borders(table)
        for r, values in enumerate([header] + [list(map(str, row)) for row in rows]):
            row = table.rows[r]
            if r == 0:
                trpr = row._tr.get_or_add_trPr()
                repeat = OxmlElement("w:tblHeader")
                repeat.set(qn("w:val"), "true")
                trpr.append(repeat)
            cant_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(cant_split)
            for c, value in enumerate(values):
                cell = row.cells[c]
                cell.width = Cm(widths_cm[c])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                self._margins(cell)
                self._shade(cell, BLUE if r == 0 else ("FFFFFF" if r % 2 else "F8FAFC"))
                paragraph = cell.paragraphs[0]
                pf = paragraph.paragraph_format
                pf.space_after, pf.space_before, pf.line_spacing = Pt(0), Pt(0), 1.15
                pf.first_line_indent = Cm(0)
                # Chain every row to the next so Word moves the whole table instead of splitting it.
                pf.keep_with_next = r < len(rows)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if (c in left_cols and r > 0) else WD_ALIGN_PARAGRAPH.CENTER
                if r == 0:
                    self.runs(paragraph, value, size=size, bold=True, color="FFFFFF")
                else:
                    self.runs(paragraph, value, size=size, bold=True if (r - 1, c) in bold_cells else None)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        spacer.paragraph_format.line_spacing = 0.6
        return table

    def save(self, path):
        self.doc.save(path)
